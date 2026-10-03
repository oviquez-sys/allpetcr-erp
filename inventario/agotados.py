"""Lista de productos agotados para decidir si se vuelven a pedir.

Pedido de Oscar (02/10/2026): la lista tiene que traer el código del
proveedor y el detalle del producto, porque con ella se arma el pedido o se
decide no comprar más. Por eso cada fila junta lo que hace falta para esa
decisión sin abrir otra pantalla: a quién se le compró la última vez, a qué
costo, cuánto se vendió en los últimos dos meses y cuántas veces se agotó.

La pantalla y el Excel salen de la MISMA función (`filas`), para que el
documento que se descarga no pueda decir algo distinto de lo que se ve.
"""
import re
from collections import defaultdict
from datetime import timedelta
from decimal import Decimal
from io import BytesIO

from django.db.models import Count, Sum
from django.utils import timezone

from .models import Agotamiento

# Ventana para "cuánto se vendió": la misma de "Qué reponer", así los dos
# reportes hablan del mismo período.
from core.reposicion import DIAS_VENTANA

# Prefijos que el ERP le agrega al código del proveedor al crear el SKU
# (comandos cargar_compra_belina / gosbi / zeedog y la compra de Special
# Care). Quitándolos queda el código tal como lo conoce el proveedor.
_PREFIJOS_INTERNOS = re.compile(r"^(BEL|GOS|ZDG|SPC)-", re.IGNORECASE)

VISTAS = {
    "pendientes": "Por decidir",
    "abiertos": "Todos los agotados",
    "historial": "Historial (ya repuestos)",
}


def codigo_proveedor(sku):
    """Código con el que se le pide el producto al proveedor.

    - "BEL-1234" → "1234" (y lo mismo con GOS-, ZDG-, SPC-).
    - Los de Buen Amigo y los del inventario inicial usan el REF del
      proveedor como SKU, así que van tal cual.
    - Los "NP…" los inventa el ERP cuando el producto entra sin código: el
      proveedor no los conoce, así que no se muestran (queda el código de
      barras, que sí sirve para pedir).
    """
    sku = (sku or "").strip()
    if not sku or sku.upper().startswith("NP"):
        return ""
    return _PREFIJOS_INTERNOS.sub("", sku)


def consulta(empresa, vista="pendientes"):
    qs = Agotamiento.objects.filter(producto__empresa=empresa).select_related(
        "producto", "producto__categoria", "movimiento", "decidido_por"
    )
    if vista == "historial":
        return qs.filter(repuesto_en__isnull=False).order_by("-repuesto_en", "-id")
    qs = qs.filter(repuesto_en__isnull=True)
    if vista == "pendientes":
        qs = qs.filter(decision=Agotamiento.Decision.PENDIENTE)
    return qs.order_by("-fecha", "-id")


def por_decidir(empresa):
    """Cuántos avisos esperan una decisión. Es lo que enciende la alarma del
    Inicio y no se cachea: un aviso que aparece dos minutos tarde no avisa."""
    return Agotamiento.objects.filter(
        producto__empresa=empresa, repuesto_en__isnull=True,
        decision=Agotamiento.Decision.PENDIENTE,
    ).count()


def _ultima_compra(producto_ids):
    """{producto_id: LineaCompra} con la compra recibida más reciente."""
    from compras.models import Compra, LineaCompra

    lineas = (
        LineaCompra.objects
        .filter(producto_id__in=producto_ids, compra__estado=Compra.Estado.RECIBIDA)
        .select_related("compra__proveedor")
        .order_by("producto_id", "-compra_id", "-id")
    )
    ultima = {}
    for linea in lineas:
        ultima.setdefault(linea.producto_id, linea)
    return ultima


def _vendidas(empresa, producto_ids, dias):
    from ventas.models import LineaVenta

    desde = timezone.localdate() - timedelta(days=dias)
    filas = (
        LineaVenta.objects
        .filter(producto_id__in=producto_ids, factura__empresa=empresa,
                factura__estado="EMI", factura__creado_en__date__gte=desde)
        .values("producto_id")
        .annotate(u=Sum("cantidad"))
    )
    return {f["producto_id"]: f["u"] or Decimal("0") for f in filas}


def filas(empresa, vista="pendientes"):
    avisos = list(consulta(empresa, vista))
    ids = [a.producto_id for a in avisos]
    compras = _ultima_compra(ids)
    vendidas = _vendidas(empresa, ids, DIAS_VENTANA)
    veces = defaultdict(int, {
        f["producto_id"]: f["n"]
        for f in Agotamiento.objects.filter(producto_id__in=ids)
        .values("producto_id").annotate(n=Count("id"))
    })
    resultado = []
    for a in avisos:
        p = a.producto
        linea = compras.get(p.pk)
        resultado.append({
            "aviso": a,
            "producto": p,
            "codigo_proveedor": codigo_proveedor(p.sku),
            "proveedor": linea.compra.proveedor.nombre if linea else "",
            "ultima_factura": (linea.compra.factura_proveedor or linea.compra.numero) if linea else "",
            "ultimo_costo": linea.costo_unitario if linea else None,
            "vendidas": vendidas.get(p.pk, Decimal("0")),
            "veces": veces[p.pk],
            "salida": a.get_tipo_salida_display() if a.tipo_salida else "",
            "referencia": a.movimiento.referencia if a.movimiento_id else "",
        })
    return resultado


COLUMNAS_EXCEL = [
    ("Código proveedor", 16), ("Código de barras", 16), ("SKU interno", 18),
    ("Producto", 46), ("Marca", 16), ("Categoría", 22), ("Proveedor (última compra)", 26),
    ("Última factura", 16), ("Último costo ₡", 14), ("Precio de venta ₡", 15),
    ("Vendidas últimos {dias} días", 14), ("Veces agotado", 10), ("Se agotó el", 13),
    ("Cómo salió", 18), ("Decisión", 16), ("Nota", 30),
]


def excel(empresa, vista="pendientes"):
    """El "documento" que pidió Oscar: la misma lista, para pedir o descartar.

    La columna Decisión va con lo que se haya marcado en pantalla; si se
    imprime para decidir en papel, los "Por decidir" quedan a la vista.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Agotados"
    hoy = timezone.localdate()
    ws.append([f"Productos agotados — {VISTAS.get(vista, '')} — {hoy:%d/%m/%Y}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    encabezados = [t.format(dias=DIAS_VENTANA) for t, _ in COLUMNAS_EXCEL]
    ws.append(encabezados)
    relleno = PatternFill("solid", fgColor="0B3161")
    for i, (_, ancho) in enumerate(COLUMNAS_EXCEL, start=1):
        celda = ws.cell(row=3, column=i)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = relleno
        celda.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = ancho
    for f in filas(empresa, vista):
        p, a = f["producto"], f["aviso"]
        ws.append([
            f["codigo_proveedor"], p.codigo_barras, p.sku, p.nombre, p.marca,
            p.categoria.nombre if p.categoria_id else "", f["proveedor"], f["ultima_factura"],
            float(f["ultimo_costo"]) if f["ultimo_costo"] is not None else None,
            float(p.precio_venta), float(f["vendidas"]), f["veces"],
            timezone.localtime(a.fecha).strftime("%d/%m/%Y"), f["salida"],
            a.get_decision_display(), a.nota,
        ])
    for fila in ws.iter_rows(min_row=4, min_col=9, max_col=10):
        for celda in fila:
            celda.number_format = "#,##0.00"
    ws.freeze_panes = "A4"
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()
