"""Qué se vendió en un rango de fechas, para reponerlo.

Pedido de Oscar (20/10/2026): "encuentro que se agotó y que quedan pocas
unidades, pero no sé qué se vendió para reponer… ojalá con un rango de
fechas". Los otros reportes no responden eso directo: "Qué reponer" usa una
ventana fija de 60 días y un cálculo de velocidad, "Más vendidos" no dice
cuánto queda ni a quién se le compra, y "Agotados" solo muestra lo que llegó
a cero.

Cada fila: un producto que salió en el rango, cuántas unidades salieron,
cuántas quedan HOY, a quién se le compró la última vez y con qué código se le
pide. Primero lo agotado, después lo que quedó con pocas, después el resto.

- Cuenta facturas EMITIDAS (las anuladas no). Las regalías cuentan, porque
  también vacían el estante. Lo que se devolvió de esas mismas ventas se
  resta: volvió al estante y no hay que reponerlo.
- "Pocas" = existencia en o bajo el mínimo del producto (el mismo criterio
  de "Stock contra el mínimo").
- La pantalla y el Excel salen de la MISMA función (`filas`).
"""
from datetime import timedelta
from decimal import Decimal
from io import BytesIO

from django.db.models import Sum
from django.utils import timezone

from .agotados import _ultima_compra, codigo_proveedor

# Más de un año de una vez no sirve para reponer y pone lenta la pantalla.
MAX_DIAS = 366

ATAJOS = [
    ("hoy", "Hoy"),
    ("ayer", "Ayer"),
    ("7", "Últimos 7 días"),
    ("15", "Últimos 15 días"),
    ("mes", "Este mes"),
]

AGOTADO, POCAS, HAY = "agotado", "pocas", "hay"
ESTADOS = {AGOTADO: "Agotado", POCAS: "Quedan pocas", HAY: "Hay"}


def rango_de_atajo(clave, hoy=None):
    hoy = hoy or timezone.localdate()
    if clave == "hoy":
        return hoy, hoy
    if clave == "ayer":
        ayer = hoy - timedelta(days=1)
        return ayer, ayer
    if clave == "mes":
        return hoy.replace(day=1), hoy
    if clave in ("7", "15"):
        return hoy - timedelta(days=int(clave) - 1), hoy
    return None


def normalizar(desde, hasta, hoy=None):
    """Rango válido: en orden, sin futuro y de máximo MAX_DIAS."""
    hoy = hoy or timezone.localdate()
    hasta = min(hasta or hoy, hoy)
    desde = desde or hasta - timedelta(days=6)
    if desde > hasta:
        desde, hasta = hasta, desde
    if (hasta - desde).days + 1 > MAX_DIAS:
        desde = hasta - timedelta(days=MAX_DIAS - 1)
    return desde, hasta


def filas(empresa, desde, hasta, proveedor=""):
    from ventas.models import LineaDevolucion, LineaVenta

    lineas = LineaVenta.objects.filter(
        factura__empresa=empresa,
        factura__estado="EMI",
        factura__creado_en__date__gte=desde,
        factura__creado_en__date__lte=hasta,
    )
    salieron = {
        f["producto_id"]: f["u"] or Decimal("0")
        for f in lineas.values("producto_id").annotate(u=Sum("cantidad"))
    }
    devueltas = {
        f["linea_venta__producto_id"]: f["u"] or Decimal("0")
        for f in LineaDevolucion.objects.filter(linea_venta__in=lineas)
        .values("linea_venta__producto_id").annotate(u=Sum("cantidad"))
    }
    netas = {pid: u - devueltas.get(pid, Decimal("0")) for pid, u in salieron.items()}
    netas = {pid: u for pid, u in netas.items() if u > 0}

    from catalogo.models import Producto

    productos = Producto.objects.filter(empresa=empresa, pk__in=netas).select_related("categoria")
    compras = _ultima_compra(list(netas))
    resultado = []
    for p in productos:
        linea = compras.get(p.pk)
        nombre_prov = linea.compra.proveedor.nombre if linea else ""
        if proveedor and nombre_prov != proveedor:
            continue
        if p.stock_actual <= 0:
            estado = AGOTADO
        elif p.stock_actual <= p.stock_minimo:
            estado = POCAS
        else:
            estado = HAY
        resultado.append({
            "producto": p,
            "vendidas": netas[p.pk],
            "quedan": p.stock_actual,
            "estado": estado,
            "estado_texto": ESTADOS[estado],
            "codigo_proveedor": codigo_proveedor(p.sku),
            "proveedor": nombre_prov,
        })
    orden = {AGOTADO: 0, POCAS: 1, HAY: 2}
    resultado.sort(key=lambda f: (orden[f["estado"]], -f["vendidas"], f["producto"].nombre))
    return resultado


def resumen(lista):
    return {
        "productos": len(lista),
        "unidades": sum((f["vendidas"] for f in lista), Decimal("0")),
        "agotados": sum(1 for f in lista if f["estado"] == AGOTADO),
        "pocas": sum(1 for f in lista if f["estado"] == POCAS),
    }


COLUMNAS_EXCEL = [
    ("Código proveedor", 16), ("Código de barras", 16), ("Producto", 46), ("Marca", 16),
    ("Categoría", 22), ("Proveedor (última compra)", 26), ("Vendidas", 10), ("Quedan hoy", 11),
    ("Estado", 14), ("A pedir", 10),
]


def excel(empresa, desde, hasta, proveedor="", solo_reponer=False):
    """La misma lista, con una columna "A pedir" vacía para anotar el pedido."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Vendido"
    titulo = f"Lo vendido del {desde:%d/%m/%Y} al {hasta:%d/%m/%Y}"
    if proveedor:
        titulo += f" — {proveedor}"
    ws.append([titulo])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append([t for t, _ in COLUMNAS_EXCEL])
    relleno = PatternFill("solid", fgColor="0B3161")
    for i, (_, ancho) in enumerate(COLUMNAS_EXCEL, start=1):
        celda = ws.cell(row=3, column=i)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = relleno
        celda.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = ancho
    for f in filas(empresa, desde, hasta, proveedor):
        if solo_reponer and f["estado"] == HAY:
            continue
        p = f["producto"]
        ws.append([
            f["codigo_proveedor"], p.codigo_barras, p.nombre, p.marca,
            p.categoria.nombre if p.categoria_id else "", f["proveedor"],
            float(f["vendidas"]), float(f["quedan"]), f["estado_texto"], None,
        ])
    ws.freeze_panes = "A4"
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()
