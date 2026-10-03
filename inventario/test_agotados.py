"""Aviso de productos agotados (pedido de Oscar, 02/10/2026).

Lo que fijan estas pruebas: todo producto que tuvo existencia y quedó en cero
por una salida abre un aviso; volver a tener existencia lo cierra; la lista
trae el código del proveedor; y la alarma se ve en el Inicio y en el POS.
"""
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import Group, User
from django.test import TestCase
from django.urls import reverse

from catalogo.models import Producto
from compras.models import Compra, LineaCompra, Proveedor
from core.models import Empresa, Sucursal
from ventas.tests import BaseVentas

from .agotados import codigo_proveedor, filas, por_decidir
from .models import Agotamiento, Bodega
from .services import registrar_movimiento


class BaseAgotados(TestCase):
    def setUp(self):
        self.empresa = Empresa.objects.create(nombre="ALLPETCR.COM")
        self.sucursal = Sucursal.objects.create(empresa=self.empresa, nombre="Central")
        self.bodega = Bodega.objects.create(sucursal=self.sucursal, nombre="Principal")
        self.gerente = User.objects.create_user("oscar", password="x", is_staff=True, is_superuser=True)
        self.producto = self.nuevo("BEL-1234", "Belina Adulto 2 kg", 3)

    def nuevo(self, sku, nombre, stock):
        p = Producto.objects.create(empresa=self.empresa, sku=sku, nombre=nombre,
                                    precio_venta=Decimal("5000"), codigo_barras="7441000000001")
        registrar_movimiento(producto=p, bodega=self.bodega, tipo="INI",
                             cantidad=Decimal(stock), costo_unitario=Decimal("2000"), referencia="INI")
        return p

    def mover(self, tipo, cantidad, producto=None):
        return registrar_movimiento(producto=producto or self.producto, bodega=self.bodega,
                                    tipo=tipo, cantidad=Decimal(cantidad), referencia=f"{tipo}-1")


class CuandoSeAbreElAviso(BaseAgotados):
    def test_la_venta_de_la_ultima_unidad_abre_el_aviso(self):
        self.mover("VEN", -2)
        self.assertFalse(Agotamiento.objects.exists(), "Todavía queda 1: no está agotado")
        mov = self.mover("VEN", -1)
        aviso = Agotamiento.objects.get()
        self.assertEqual(aviso.producto, self.producto)
        self.assertEqual(aviso.movimiento, mov)
        self.assertEqual(aviso.tipo_salida, "VEN")
        self.assertEqual(aviso.decision, Agotamiento.Decision.PENDIENTE)

    def test_regalia_y_ajuste_tambien_avisan(self):
        otro = self.nuevo("85313", "Juguete", 1)
        self.mover("REG", -3)
        self.mover("AJU", -1, producto=otro)
        self.assertEqual(
            dict(Agotamiento.objects.values_list("producto__sku", "tipo_salida")),
            {"BEL-1234": "REG", "85313": "AJU"},
        )

    def test_anular_una_compra_no_avisa(self):
        # Corrige una compra mal digitada: no es un producto que se acabó.
        self.mover("DEV", -3)
        self.assertFalse(Agotamiento.objects.exists())

    def test_un_solo_aviso_abierto_por_producto(self):
        self.mover("VEN", -3)
        self.mover("DEV", 1)   # anulan la venta: vuelve 1
        self.mover("VEN", -1)  # se vende otra vez
        self.assertEqual(Agotamiento.objects.filter(repuesto_en__isnull=True).count(), 1)
        self.assertEqual(Agotamiento.objects.count(), 2, "El primero queda en el historial")


class CuandoSeCierra(BaseAgotados):
    def test_la_compra_cierra_el_aviso(self):
        self.mover("VEN", -3)
        self.mover("COM", 6)
        aviso = Agotamiento.objects.get()
        self.assertIsNotNone(aviso.repuesto_en)
        self.assertEqual(por_decidir(self.empresa), 0)

    def test_la_salida_que_no_agota_no_toca_el_aviso(self):
        self.mover("VEN", -1)
        self.assertFalse(Agotamiento.objects.exists())


class CodigoDelProveedor(TestCase):
    def test_quita_los_prefijos_internos(self):
        self.assertEqual(codigo_proveedor("BEL-1234"), "1234")
        self.assertEqual(codigo_proveedor("SPC-MPET017"), "MPET017")
        self.assertEqual(codigo_proveedor("ZDG-7898000000000"), "7898000000000")
        self.assertEqual(codigo_proveedor("GOS-77"), "77")

    def test_ref_de_buen_amigo_va_tal_cual(self):
        self.assertEqual(codigo_proveedor("85313"), "85313")

    def test_el_codigo_inventado_por_el_erp_no_se_muestra(self):
        self.assertEqual(codigo_proveedor("NP261002101010123456"), "")


class ListaYDocumento(BaseAgotados):
    def setUp(self):
        super().setUp()
        proveedor = Proveedor.objects.create(empresa=self.empresa, nombre="Distribuidora Belina")
        compra = Compra.objects.create(empresa=self.empresa, sucursal=self.sucursal, proveedor=proveedor,
                                       numero="C-1", factura_proveedor="F-777",
                                       estado=Compra.Estado.RECIBIDA)
        LineaCompra.objects.create(compra=compra, producto=self.producto, cantidad=Decimal("3"),
                                   costo_unitario=Decimal("2100"), total=Decimal("6300"))
        self.mover("VEN", -3)
        self.client.login(username="oscar", password="x")

    def test_la_fila_trae_lo_necesario_para_pedir(self):
        f = filas(self.empresa)[0]
        self.assertEqual(f["codigo_proveedor"], "1234")
        self.assertEqual(f["proveedor"], "Distribuidora Belina")
        self.assertEqual(f["ultima_factura"], "F-777")
        self.assertEqual(f["ultimo_costo"], Decimal("2100"))
        self.assertEqual(f["salida"], "Venta")

    def test_pantalla_muestra_el_agotado(self):
        r = self.client.get(reverse("inventario:agotados"))
        self.assertContains(r, "Belina Adulto 2 kg")
        self.assertContains(r, "1234")

    def test_decidir_lo_saca_de_por_decidir(self):
        aviso = Agotamiento.objects.get()
        r = self.client.post(reverse("inventario:agotados") + "?ver=pendientes",
                             {"aviso": aviso.pk, "decision": "NO", "nota": "Se vende poco"})
        self.assertEqual(r.status_code, 302)
        aviso.refresh_from_db()
        self.assertEqual(aviso.decision, "NO")
        self.assertEqual(aviso.nota, "Se vende poco")
        self.assertEqual(aviso.decidido_por, self.gerente)
        self.assertEqual(por_decidir(self.empresa), 0)
        r = self.client.get(reverse("inventario:agotados") + "?ver=abiertos")
        self.assertContains(r, "Belina Adulto 2 kg", msg_prefix="Sigue agotado: aparece en 'Todos'")

    def test_decision_invalida_no_se_guarda(self):
        aviso = Agotamiento.objects.get()
        self.client.post(reverse("inventario:agotados"), {"aviso": aviso.pk, "decision": "XX"})
        aviso.refresh_from_db()
        self.assertEqual(aviso.decision, "PEN")

    def test_excel_con_codigo_del_proveedor(self):
        from openpyxl import load_workbook

        r = self.client.get(reverse("inventario:agotados_excel"))
        self.assertEqual(r.status_code, 200)
        ws = load_workbook(BytesIO(r.content)).active
        self.assertEqual(ws.cell(row=3, column=1).value, "Código proveedor")
        fila = [c.value for c in ws[4]]
        self.assertEqual(fila[0], "1234")
        self.assertEqual(fila[3], "Belina Adulto 2 kg")
        self.assertEqual(fila[6], "Distribuidora Belina")
        self.assertEqual(fila[14], "Por decidir")

    def test_el_cajero_no_entra_a_la_lista(self):
        cajero = User.objects.create_user("caja", password="x", is_staff=True)
        cajero.groups.add(Group.objects.get_or_create(name="Cajero")[0])
        self.client.login(username="caja", password="x")
        r = self.client.get(reverse("inventario:agotados"))
        self.assertNotEqual(r.status_code, 200)
        r = self.client.get(reverse("inventario:agotados_excel"))
        self.assertNotEqual(r.status_code, 200)

    def test_alarma_en_el_inicio(self):
        r = self.client.get("/")
        self.assertContains(r, "1 producto se agotó")
        self.assertContains(r, reverse("inventario:agotados"))


class AlarmaEnElPOS(BaseVentas):
    def test_la_venta_que_agota_lo_avisa_en_la_caja(self):
        self.client.login(username="oscar", password="clave-test")
        r = self.client.post(
            reverse("ventas:vender"),
            data='{"medio_pago":"EFE","lineas":[{"producto_id":%d,"cantidad":10}]}' % self.producto.pk,
            content_type="application/json",
        )
        self.assertEqual(r.json()["agotados"], ["Arnés prueba"])
        self.assertTrue(Agotamiento.objects.filter(producto=self.producto).exists())

    def test_la_venta_que_no_agota_no_avisa(self):
        self.client.login(username="oscar", password="clave-test")
        r = self.client.post(
            reverse("ventas:vender"),
            data='{"medio_pago":"EFE","lineas":[{"producto_id":%d,"cantidad":1}]}' % self.producto.pk,
            content_type="application/json",
        )
        self.assertEqual(r.json()["agotados"], [])


class AgotadosDeAntes(BaseAgotados):
    """La migración crea el aviso para lo que ya estaba agotado el día que
    nació la función (si no, esos seguirían escondidos)."""

    def test_crea_el_aviso_solo_para_los_que_salieron(self):
        from importlib import import_module

        from django.apps import apps

        crear = import_module("inventario.migrations.0007_agotamientos_existentes").crear
        anulado = self.nuevo("GOS-9", "Compra anulada", 2)
        nunca = Producto.objects.create(empresa=self.empresa, sku="X-0", nombre="Nunca tuvo")
        self.mover("VEN", -3)
        self.mover("DEV", -2, producto=anulado)
        Agotamiento.objects.all().delete()  # como si la función no existiera aún

        crear(apps, None)
        crear(apps, None)  # correrla dos veces no duplica

        self.assertEqual(list(Agotamiento.objects.values_list("producto__sku", flat=True)), ["BEL-1234"])
        self.assertFalse(Agotamiento.objects.filter(producto=nunca).exists())
