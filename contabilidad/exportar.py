"""Excel para el contador (auditoría 26/09/2026, CAJ-01).

Hasta hoy el contador solo podía mirar el ERP en pantalla: para trabajar con
los números tenía que copiarlos a mano. Este archivo junta en un solo Excel,
para el rango de fechas que se pida, todo lo que un contador suele necesitar:

- Ventas: una fila por tiquete, con cliente, cédula, medio de pago, subtotal,
  descuento, IVA y total (las anuladas se incluyen y se marcan).
- Detalle de ventas: una fila por producto vendido, con tarifa, CABYS, IVA y
  costo (las ventas anteriores al 26/09/2026 no guardaban tarifa por línea:
  esas celdas salen vacías, no inventadas).
- Pagos: cómo entró el dinero de cada venta, con los pagos mixtos repartidos.
- Devoluciones.
- Compras: con proveedor, factura, IVA acreditable y total.
- Libro diario: todas las líneas de asiento.

Los montos van como números (no texto) para que se puedan sumar en Excel.
"""
import io
from datetime import datetime, time

from django.utils import timezone


def _rango(desde, hasta):
    zona = timezone.get_current_timezone()
    return (timezone.make_aware(datetime.combine(desde, time.min), zona),
            timezone.make_aware(datetime.combine(hasta, time.max), zona))


def _hoja(wb, titulo, encabezado, filas, anchos=None):
    from openpyxl.styles import Font, PatternFill

    ws = wb.create_sheet(titulo)
    ws.append(encabezado)
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="0B3161")
    for fila in filas:
        ws.append(fila)
    ws.freeze_panes = "A2"
    for i, ancho in enumerate(anchos or [], start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = ancho
    return ws


def _n(valor):
    return float(valor) if valor is not None else None


def _local(momento):
    return timezone.localtime(momento).replace(tzinfo=None) if momento else None


def excel_contador(empresa, desde, hasta) -> bytes:
    from openpyxl import Workbook

    from compras.models import Compra
    from ventas.models import DevolucionVenta, FacturaVenta, LineaVenta

    from .models import LineaAsiento

    ini, fin = _rango(desde, hasta)
    wb = Workbook()
    wb.remove(wb.active)

    facturas = (FacturaVenta.objects.filter(empresa=empresa, creado_en__range=(ini, fin))
                .select_related("cliente", "usuario").prefetch_related("pagos").order_by("creado_en", "id"))
    _hoja(wb, "Ventas", ["Número", "Fecha", "Cliente", "Cédula", "Medio de pago", "Subtotal (sin IVA)",
                         "Descuento", "IVA", "Total", "Estado", "Motivo anulación", "Atendió"],
          [[f.numero, _local(f.creado_en), f.cliente.nombre if f.cliente else "Consumidor final",
            f.cliente.identificacion if f.cliente else "", f.get_medio_pago_display(), _n(f.subtotal),
            _n(f.descuento), _n(f.impuesto), _n(f.total), f.get_estado_display(), f.motivo_anulacion,
            f.usuario.username if f.usuario else ""] for f in facturas],
          [16, 18, 26, 14, 16, 16, 12, 12, 12, 10, 26, 12])

    lineas = (LineaVenta.objects.filter(factura__in=facturas).select_related("factura", "producto")
              .order_by("factura__creado_en", "id"))
    _hoja(wb, "Detalle de ventas", ["Número", "Fecha", "Estado", "Código", "Producto", "CABYS", "Cantidad",
                                    "Precio unitario", "Descuento", "Regalía", "Tarifa IVA %",
                                    "Subtotal (sin IVA)", "IVA", "Total", "Costo unitario"],
          [[l.factura.numero, _local(l.factura.creado_en), l.factura.get_estado_display(), l.producto.sku,
            l.producto.nombre, l.cabys or l.producto.cabys, _n(l.cantidad), _n(l.precio_unitario),
            _n(l.descuento_monto), "Sí" if l.es_regalia else "", _n(l.tarifa_iva), _n(l.subtotal),
            _n(l.impuesto), _n(l.total), _n(l.costo_unitario)] for l in lineas],
          [16, 18, 10, 14, 36, 16, 10, 14, 12, 8, 10, 16, 12, 12, 14])

    nombres = dict(FacturaVenta.MedioPago.choices)
    pagos = []
    for f in facturas:
        if f.estado != FacturaVenta.Estado.EMITIDA:
            continue
        for medio, monto in f.desglose_pagos():
            pagos.append([f.numero, _local(f.creado_en), nombres.get(medio, medio), _n(monto)])
    _hoja(wb, "Pagos", ["Número", "Fecha", "Medio", "Monto"], pagos, [16, 18, 16, 14])

    devoluciones = (DevolucionVenta.objects.filter(factura__empresa=empresa, creado_en__range=(ini, fin))
                    .select_related("factura", "usuario").order_by("creado_en"))
    _hoja(wb, "Devoluciones", ["Número", "Fecha", "Venta original", "Medio de la venta", "Monto devuelto",
                               "Motivo", "Hecha por"],
          [[d.numero, _local(d.creado_en), d.factura.numero, d.factura.get_medio_pago_display(), _n(d.total),
            d.motivo, d.usuario.username if d.usuario else ""] for d in devoluciones],
          [16, 18, 16, 16, 14, 30, 12])

    compras = (Compra.objects.filter(empresa=empresa, creado_en__range=(ini, fin))
               .select_related("proveedor").order_by("creado_en"))
    _hoja(wb, "Compras", ["Número", "Fecha", "Proveedor", "Cédula proveedor", "Factura proveedor",
                          "Forma de pago", "Subtotal (sin IVA)", "IVA acreditable", "Total factura", "Estado"],
          [[c.numero, _local(c.recibida_en or c.creado_en), c.proveedor.nombre, c.proveedor.identificacion,
            c.factura_proveedor, c.get_forma_pago_display(), _n(c.total), _n(c.iva), _n(c.total_factura),
            c.get_estado_display()] for c in compras],
          [14, 18, 28, 16, 18, 22, 16, 14, 14, 10])

    diario = (LineaAsiento.objects.filter(asiento__empresa=empresa, asiento__fecha__range=(desde, hasta))
              .select_related("asiento", "cuenta").order_by("asiento__fecha", "asiento__id", "id"))
    _hoja(wb, "Libro diario", ["Asiento", "Fecha", "Descripción", "Origen", "Referencia", "Cuenta",
                               "Nombre de la cuenta", "Debe", "Haber", "Detalle"],
          [[l.asiento.numero, l.asiento.fecha, l.asiento.descripcion, l.asiento.get_origen_display(),
            l.asiento.referencia, l.cuenta.codigo, l.cuenta.nombre, _n(l.debe), _n(l.haber), l.detalle]
           for l in diario],
          [14, 12, 36, 14, 16, 10, 30, 14, 14, 20])

    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()
