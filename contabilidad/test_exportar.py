"""CAJ-01 y CAJ-04 (auditoría 26/09/2026): Excel para el contador y reportes de
ventas. La prueba central es que los números del Excel CUADREN entre sí:
ventas = pagos, y el libro diario balanceado."""
import io
from datetime import timedelta
from decimal import Decimal

from django.contrib.auth.models import Group, User
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from caja.services import abrir_caja
from catalogo.models import Categoria, Impuesto, Producto
from core import reportes as rep
from core.models import Empresa, Sucursal
from django.test import TestCase
from inventario.models import Bodega
from inventario.services import registrar_movimiento
from ventas.devoluciones import registrar_devolucion
from ventas.services import anular_factura, registrar_venta


class Base(TestCase):
    def setUp(self):
        self.empresa = Empresa.objects.create(nombre="ALLPETCR.COM", regimen=Empresa.Regimen.TRADICIONAL)
        sucursal = Sucursal.objects.create(empresa=self.empresa, nombre="Central")
        bodega = Bodega.objects.create(sucursal=sucursal, nombre="Principal", principal=True)
        self.usuario = User.objects.create_user("oscar", password="x", is_staff=True, is_superuser=True)
        iva = Impuesto.objects.create(nombre="IVA", tarifa=Decimal("13"))
        paseo = Categoria.objects.create(nombre="Paseo")
        arneses = Categoria.objects.create(nombre="Arneses", padre=paseo)
        self.p1 = Producto.objects.create(empresa=self.empresa, sku="P1", nombre="Arnés", precio_venta=Decimal("5650"),
                                          impuesto=iva, categoria=arneses)
        self.p2 = Producto.objects.create(empresa=self.empresa, sku="P2", nombre="Pelota", precio_venta=Decimal("1130"),
                                          impuesto=iva)
        self.dormido = Producto.objects.create(empresa=self.empresa, sku="P3", nombre="Cama", precio_venta=20000)
        for p in (self.p1, self.p2, self.dormido):
            registrar_movimiento(producto=p, bodega=bodega, tipo="INI", cantidad=Decimal("20"),
                                 costo_unitario=Decimal("1000"), referencia="INI")
        self.sesion = abrir_caja(sucursal=sucursal, usuario=self.usuario, monto_apertura=Decimal("0"))

        def vender(medio, producto, cantidad, **extra):
            return registrar_venta(sesion_caja=self.sesion, medio_pago=medio, usuario=self.usuario,
                                   lineas=[{"producto_id": producto.pk, "cantidad": cantidad}], **extra)
        self.f1 = vender("EFE", self.p1, 2)
        self.f2 = vender("MIX", self.p2, 3, pagos=[{"medio": "EFE", "monto": 1390}, {"medio": "SIN", "monto": 2000}])
        self.f3 = vender("TAR", self.p1, 1)
        anular_factura(factura=self.f3, motivo="error", usuario=self.usuario)
        registrar_devolucion(factura=self.f1, lineas=[{"linea_venta_id": self.f1.lineas.get().pk, "cantidad": 1}],
                             motivo="cambio", usuario=self.usuario)
        self.hoy = timezone.localdate()


class ExcelDelContador(Base):
    def descargar(self, usuario=None):
        self.client.force_login(usuario or self.usuario)
        r = self.client.get(reverse("contabilidad:exportar"),
                            {"descargar": "1", "desde": self.hoy.isoformat(), "hasta": self.hoy.isoformat()})
        return r, load_workbook(io.BytesIO(r.content), data_only=True)

    def test_trae_todas_las_hojas(self):
        _, wb = self.descargar()
        self.assertEqual(wb.sheetnames, ["Ventas", "Detalle de ventas", "Pagos", "Devoluciones", "Compras", "Libro diario"])

    def test_los_numeros_cuadran(self):
        _, wb = self.descargar()
        ventas = [f for f in wb["Ventas"].iter_rows(min_row=2, values_only=True)]
        self.assertEqual(len(ventas), 3)                              # la anulada se incluye y se marca
        vigentes = sum(f[8] for f in ventas if f[9] == "Emitida")
        pagos = sum(f[3] for f in wb["Pagos"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(vigentes, pagos)                             # ventas = pagos
        self.assertEqual(vigentes, 11300 + 3390)
        diario = list(wb["Libro diario"].iter_rows(min_row=2, values_only=True))
        self.assertAlmostEqual(sum(f[7] or 0 for f in diario), sum(f[8] or 0 for f in diario), places=2)
        detalle = list(wb["Detalle de ventas"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(detalle[0][10], 13)                          # tarifa por línea
        self.assertEqual(len(list(wb["Devoluciones"].iter_rows(min_row=2))), 1)

    def test_el_contador_puede_y_el_cajero_no(self):
        contador = User.objects.create_user("conta", password="x", is_staff=True)
        contador.groups.add(Group.objects.get(name="Contador"))
        r, _ = self.descargar(contador)
        self.assertEqual(r.status_code, 200)
        cajero = User.objects.create_user("caj", password="x", is_staff=True)
        cajero.groups.add(Group.objects.get(name="Cajero"))
        self.client.force_login(cajero)
        r = self.client.get(reverse("contabilidad:exportar"), {"descargar": "1"})
        self.assertEqual(r.status_code, 302)


class ReporteDeVentas(Base):
    def test_totales_y_utilidad(self):
        datos = rep.ventas_por_periodo(self.empresa, self.hoy, self.hoy, "dia")
        t = datos["totales"]
        self.assertEqual(t["tiquetes"], 2)
        self.assertEqual(t["total"], Decimal("14690"))
        self.assertEqual(t["sin_iva"] + t["iva"], t["total"])
        self.assertEqual(t["costo"], Decimal("5000"))                 # 2 arneses + 3 pelotas a ₡1.000
        self.assertEqual(t["utilidad"], t["sin_iva"] - t["costo"])
        self.assertEqual(datos["devoluciones"], Decimal("5650"))

    def test_por_categoria_usa_la_raiz(self):
        cats = {c["cat"]: c for c in rep.ventas_por_periodo(self.empresa, self.hoy, self.hoy)["por_categoria"]}
        self.assertIn("Paseo", cats)
        self.assertIn("Sin categoría", cats)

    def test_menos_vendidos_empieza_por_lo_dormido(self):
        primeros = rep.menos_vendidos(self.empresa, self.hoy - timedelta(days=30), self.hoy)
        self.assertEqual(primeros[0], self.dormido)
        self.assertEqual(primeros[0].vendidas, 0)

    def test_la_pantalla_abre(self):
        self.client.force_login(self.usuario)
        for agrupar in ("dia", "semana", "mes"):
            r = self.client.get(reverse("core:reporte_ventas"), {"agrupar": agrupar})
            self.assertEqual(r.status_code, 200)
