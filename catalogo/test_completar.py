"""Pruebas de la pantalla "Completar catálogo" (catalogo/completar.py)."""
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from core.models import AuditLog, Empresa

from . import completar
from .models import Categoria, Producto


def _excel(filas, encabezado=("Código (SKU)", "Nombre", "Existencia", "Mascota", "Categoría", "Descripción")):
    wb = Workbook()
    ws = wb.active
    ws.title = "Productos"
    ws.append(list(encabezado))
    for f in filas:
        ws.append(list(f))
    b = BytesIO()
    wb.save(b)
    return SimpleUploadedFile("x.xlsx", b.getvalue())


class CompletarCatalogo(TestCase):
    def setUp(self):
        self.empresa = Empresa.objects.create(nombre="ALLPETCR.COM")
        self.juguetes = Categoria.objects.create(nombre="Juguetes")
        self.pelotas = Categoria.objects.create(nombre="Pelotas", padre=self.juguetes)
        self.higiene = Categoria.objects.create(nombre="Higiene")
        self.completo = self._p("1", "Completo", mascota="Perro", categoria=self.pelotas, descripcion="Ya está.")
        self.sin_mascota = self._p("2", "Sin mascota", categoria=self.pelotas, descripcion="Texto")
        self.mascota_rara = self._p("3", "Mascota rara", mascota="perros", categoria=self.pelotas, descripcion="Texto")
        self.sin_cat = self._p("4", "Sin categoría", mascota="Gato", descripcion="Texto")
        self.sin_desc = self._p("5", "Sin descripción", mascota="Gato", categoria=self.higiene, descripcion="   ")
        self.agotado = self._p("6", "Agotado", stock=0)
        User.objects.create_user("oscar", password="x", is_staff=True, is_superuser=True)
        self.client.login(username="oscar", password="x")

    def _p(self, sku, nombre, mascota="", categoria=None, descripcion="", stock=5):
        return Producto.objects.create(
            empresa=self.empresa, sku=sku, nombre=nombre, mascota=mascota, categoria=categoria,
            descripcion=descripcion, precio_venta=Decimal("1000"), stock_actual=Decimal(stock),
        )

    # --- qué cuenta como pendiente -------------------------------------
    def test_pendientes_incluye_mascota_que_la_web_no_reconoce(self):
        skus = set(completar.pendientes(self.empresa).values_list("sku", flat=True))
        self.assertEqual(skus, {"2", "3", "4", "5"})

    def test_agotados_solo_si_se_piden(self):
        self.assertNotIn("6", completar.pendientes(self.empresa).values_list("sku", flat=True))
        self.assertIn("6", completar.pendientes(self.empresa, incluir_agotados=True).values_list("sku", flat=True))

    def test_conteos_por_campo(self):
        c = completar.conteos(self.empresa)
        self.assertEqual((c["mascota"], c["categoria"], c["descripcion"], c["total"]), (2, 1, 1, 4))

    def test_normalizar_mascota(self):
        self.assertEqual(completar.normalizar_mascota("  PERROS "), "Perro")
        self.assertEqual(completar.normalizar_mascota("Ambos"), "Perro y gato")
        self.assertEqual(completar.normalizar_mascota("Felinos"), "Gato")
        self.assertIsNone(completar.normalizar_mascota("dinosaurio"))

    # --- pantalla --------------------------------------------------------
    def test_pantalla_lista_pendientes(self):
        r = self.client.get(reverse("catalogo:completar"))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Sin mascota")
        self.assertNotContains(r, ">Completo<")

    def test_guardar_desde_pantalla_y_vacio_no_borra(self):
        r = self.client.post(reverse("catalogo:completar"), {
            "producto": [self.sin_mascota.pk, self.sin_cat.pk],
            f"m_{self.sin_mascota.pk}": "Perro y gato",
            f"c_{self.sin_mascota.pk}": "",   # vacío: conserva Pelotas
            f"d_{self.sin_mascota.pk}": "",   # vacío: conserva "Texto"
            f"m_{self.sin_cat.pk}": "Gato",
            f"c_{self.sin_cat.pk}": str(self.higiene.pk),
            f"d_{self.sin_cat.pk}": "Texto",
        })
        self.assertEqual(r.status_code, 302)
        self.sin_mascota.refresh_from_db()
        self.sin_cat.refresh_from_db()
        self.assertEqual(self.sin_mascota.mascota, "Perro y gato")
        self.assertEqual(self.sin_mascota.categoria, self.pelotas)
        self.assertEqual(self.sin_mascota.descripcion, "Texto")
        self.assertEqual(self.sin_cat.categoria, self.higiene)

    def test_cambio_queda_en_auditoria(self):
        antes = AuditLog.objects.filter(tabla="catalogo.producto").count()
        self.client.post(reverse("catalogo:completar"), {
            "producto": [self.sin_mascota.pk], f"m_{self.sin_mascota.pk}": "Perro",
        })
        self.assertGreater(AuditLog.objects.filter(tabla="catalogo.producto").count(), antes)

    def test_cajero_no_entra(self):
        User.objects.create_user("caja", password="x", is_staff=True)
        self.client.login(username="caja", password="x")
        r = self.client.get(reverse("catalogo:completar"))
        self.assertEqual(r.status_code, 302)

    # --- Excel -----------------------------------------------------------
    def test_descarga_trae_pendientes_y_listas(self):
        r = self.client.get(reverse("catalogo:completar_excel"))
        wb = load_workbook(BytesIO(r.content))
        skus = [row[0] for row in wb["Productos"].iter_rows(min_row=2, values_only=True)]
        self.assertEqual(sorted(skus), ["2", "3", "4", "5"])
        opciones = [row[1] for row in wb["Opciones"].iter_rows(min_row=2, values_only=True) if row[1]]
        self.assertIn("Juguetes › Pelotas", opciones)

    def test_ida_y_vuelta_con_el_excel_descargado(self):
        """Lo que descarga la pantalla se puede volver a subir tal cual."""
        r = self.client.get(reverse("catalogo:completar_excel"))
        wb = load_workbook(BytesIO(r.content))
        ws = wb["Productos"]
        for fila in ws.iter_rows(min_row=2):
            if fila[0].value == "2":
                fila[3].value = "Perro"
        b = BytesIO()
        wb.save(b)
        rev = completar.revisar_excel(self.empresa, BytesIO(b.getvalue()))
        # El "3" tenía "perros" (invisible en la web): al volver a subir el
        # archivo sin tocarlo, la variante se corrige sola a "Perro" — y la
        # vista previa lo muestra como cambio, así que no es silencioso.
        self.assertEqual(sorted((c.sku, c.campo, c.despues) for c in rev.cambios),
                         [("2", "mascota", "Perro"), ("3", "mascota", "Perro")])
        self.assertEqual(rev.errores, [])

    def test_subir_no_guarda_hasta_confirmar(self):
        archivo = _excel([("2", "", "", "Perro", "", "Pelota de goma resistente.")])
        r = self.client.post(reverse("catalogo:completar_subir"), {"archivo": archivo})
        self.assertEqual(r.status_code, 200)
        self.sin_mascota.refresh_from_db()
        self.assertEqual(self.sin_mascota.mascota, "")  # todavía nada
        self.client.post(reverse("catalogo:completar_confirmar"), {"firma": r.context["firma"]})
        self.sin_mascota.refresh_from_db()
        self.assertEqual(self.sin_mascota.mascota, "Perro")
        self.assertEqual(self.sin_mascota.descripcion, "Pelota de goma resistente.")

    def test_firma_alterada_no_aplica(self):
        r = self.client.post(reverse("catalogo:completar_confirmar"), {"firma": "basura"})
        self.assertEqual(r.status_code, 302)
        self.sin_mascota.refresh_from_db()
        self.assertEqual(self.sin_mascota.mascota, "")

    def test_errores_no_frenan_el_resto_de_la_fila(self):
        rev = completar.revisar_excel(self.empresa, _excel([
            ("2", "", "", "dinosaurio", "Juguetes › Pelotas", "Descripción nueva"),
            ("4", "", "", "", "Categoría inventada", ""),
            ("999", "", "", "Perro", "", ""),
        ]))
        campos = {(c.sku, c.campo) for c in rev.cambios}
        self.assertEqual(campos, {("2", "descripcion")})  # la categoría ya era Pelotas
        self.assertEqual(len(rev.errores), 3)

    def test_categoria_con_raiz_equivocada_se_rechaza(self):
        self.assertIsNone(completar.buscar_categoria("Higiene › Pelotas"))
        self.assertEqual(completar.buscar_categoria("juguetes > pelotas"), self.pelotas)
        self.assertEqual(completar.buscar_categoria("Pelotas"), self.pelotas)

    def test_celdas_vacias_no_borran(self):
        rev = completar.revisar_excel(self.empresa, _excel([("1", "", "", "", "", "")]))
        self.assertEqual(rev.cambios, [])
        self.assertEqual(rev.sin_cambio, 1)

    def test_sku_numerico_de_excel(self):
        """Excel convierte 00065 o 86955 en número; 86955.0 debe leerse como "86955"."""
        self._p("86955", "Lámpara solar")
        rev = completar.revisar_excel(self.empresa, _excel([(86955, "", "", "Otros", "", "")]))
        self.assertEqual([(c.sku, c.despues) for c in rev.cambios], [("86955", "Otros")])
