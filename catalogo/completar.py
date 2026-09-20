"""Completar el catálogo: productos a los que les falta mascota, categoría o
descripción (20/09/2026).

POR QUÉ EXISTE
--------------
La carga de la compra del 19/09 dejó productos con datos incompletos: el
Excel del proveedor traía la mascota vacía en 71 filas y la descripción
vacía en 80, y dos productos quedaron sin categoría a propósito. No es un
detalle cosmético:

- **Sin mascota** el producto no aparece en los menús "Perros" y "Gatos"
  del sitio. El sitio compara el texto EXACTO (`allpetcr-web/lib/
  navegacion.ts::ESPECIES`): "Perro", "Gato" o "Perro y gato". Un "perros"
  o un "Ambos" escrito a mano es tan invisible como un campo vacío, y por
  eso esta pantalla también marca como pendiente cualquier valor que no
  sea uno de la lista.
- **Sin categoría** no cae en ninguna rama del menú.
- **Sin descripción** la ficha del producto sale vacía.

Y va a volver a pasar con cada compra nueva. Por eso es una pantalla del
ERP y no un script de una sola vez: la usa quien esté en la tienda (o
Francisco desde donde esté) sin contraseñas de base de datos ni `.bat`.

DOS FORMAS DE LLENAR
--------------------
1. En pantalla, producto por producto.
2. Por Excel: se descargan los pendientes, se llenan (a mano o con ayuda)
   y se suben. Antes de aplicar, el ERP muestra CADA cambio —valor de antes
   y valor nuevo— y nada se guarda hasta que el gerente confirma.

LA REGLA QUE NO SE NEGOCIA: UNA CELDA VACÍA NUNCA BORRA
-------------------------------------------------------
Es la misma regla de `cargar_compra_19_09`: si el Excel o el formulario
traen un campo vacío, se deja lo que había. Borrar datos hechos a mano por
una celda que alguien no llenó es el error que no se nota hasta que un
cliente no encuentra el producto. Para borrar de verdad está el admin.

Los cambios quedan en la bitácora de auditoría sin hacer nada especial:
`catalogo.producto` está en `core.signals.AUDITED`.
"""
import unicodedata
from dataclasses import dataclass, field
from io import BytesIO

from django.db import transaction
from django.db.models import Q

from .consultas import productos_visibles
from .models import Categoria, Producto

# Valores válidos, en el orden en que aparecen en la lista desplegable.
# Los tres primeros son los que entiende el menú del sitio; los otros
# existen en el modelo para lo que no es de perro ni de gato.
MASCOTAS = ["Perro", "Gato", "Perro y gato", "Peces", "Tortugas", "Otros"]

# Variantes que la gente escribe de verdad → valor oficial. Se compara sin
# mayúsculas ni tildes ni espacios de más.
_SINONIMOS = {
    "perro": "Perro", "perros": "Perro", "canino": "Perro", "caninos": "Perro",
    "gato": "Gato", "gatos": "Gato", "felino": "Gato", "felinos": "Gato",
    "perro y gato": "Perro y gato", "perros y gatos": "Perro y gato",
    "perro/gato": "Perro y gato", "perro gato": "Perro y gato",
    "ambos": "Perro y gato", "gato y perro": "Perro y gato",
    "perro-gato": "Perro y gato",
    "pez": "Peces", "peces": "Peces",
    "tortuga": "Tortugas", "tortugas": "Tortugas",
    "otro": "Otros", "otros": "Otros",
}

CAMPOS = ("mascota", "categoria", "descripcion")
ETIQUETA_CAMPO = {"mascota": "Mascota", "categoria": "Categoría", "descripcion": "Descripción"}
HOJA_PRODUCTOS = "Productos"
HOJA_OPCIONES = "Opciones"
SEPARADOR = " › "


def _plano(texto) -> str:
    texto = " ".join(str(texto or "").split()).lower()
    return "".join(c for c in unicodedata.normalize("NFD", texto) if unicodedata.category(c) != "Mn")


def normalizar_mascota(valor) -> str | None:
    """Valor oficial para lo que escribió alguien, o None si no se reconoce."""
    return _SINONIMOS.get(_plano(valor))


def etiqueta_categoria(cat) -> str:
    if cat is None:
        return ""
    return f"{cat.padre.nombre}{SEPARADOR}{cat.nombre}" if cat.padre_id else cat.nombre


def opciones_categoria():
    """[(id, etiqueta)] ordenado como el menú: raíz y debajo sus hijas."""
    todas = list(Categoria.objects.select_related("padre"))
    raices = [c for c in todas if c.padre_id is None]
    salida = []
    for r in raices:  # ya vienen por (orden, nombre)
        salida.append((r.id, r.nombre))
        salida.extend((h.id, etiqueta_categoria(h)) for h in todas if h.padre_id == r.id)
    return salida


def buscar_categoria(texto):
    """Categoria por su etiqueta ("Juguetes › Pelotas") o por su nombre solo.

    `Categoria.nombre` es único, así que el último tramo alcanza para
    identificarla; si además trae la raíz, se exige que calce, para no
    aceptar en silencio una subcategoría colgada de otra rama."""
    partes = [p.strip() for p in str(texto or "").replace(">", "›").split("›") if p.strip()]
    if not partes:
        return None
    por_nombre = {_plano(c.nombre): c for c in Categoria.objects.select_related("padre")}
    cat = por_nombre.get(_plano(partes[-1]))
    if cat is None:
        return None
    if len(partes) > 1 and (cat.padre is None or _plano(cat.padre.nombre) != _plano(partes[-2])):
        return None
    return cat


def faltantes(producto) -> list[str]:
    falta = []
    if producto.mascota not in MASCOTAS:
        falta.append("mascota")
    if producto.categoria_id is None:
        falta.append("categoria")
    if not (producto.descripcion or "").strip():
        falta.append("descripcion")
    return falta


def _filtro_falta(campo):
    if campo == "mascota":
        return ~Q(mascota__in=MASCOTAS)
    if campo == "categoria":
        return Q(categoria__isnull=True)
    return Q(descripcion__isnull=True) | Q(descripcion__regex=r"^\s*$")


def pendientes(empresa, *, falta=None, incluir_agotados=False):
    """Productos a los que les falta algo. `falta` limita a un solo campo."""
    base = productos_visibles(empresa, incluir_agotados=incluir_agotados)
    campos = [falta] if falta in CAMPOS else list(CAMPOS)
    condicion = Q()
    for c in campos:
        condicion |= _filtro_falta(c)
    return base.filter(condicion).select_related("categoria", "categoria__padre").order_by("nombre")


def conteos(empresa, *, incluir_agotados=False):
    base = productos_visibles(empresa, incluir_agotados=incluir_agotados)
    datos = {c: base.filter(_filtro_falta(c)).count() for c in CAMPOS}
    datos["total"] = pendientes(empresa, incluir_agotados=incluir_agotados).count()
    datos["catalogo"] = base.count()
    return datos


# --------------------------------------------------------------------------
# Cambios: se calculan primero (para mostrarlos) y se aplican después.
# --------------------------------------------------------------------------

@dataclass
class Cambio:
    sku: str
    nombre: str
    campo: str
    antes: str
    despues: str
    categoria_id: int | None = None  # solo para campo == "categoria"
    producto: object = None  # para mostrar la foto; no viaja en la firma

    def a_dict(self):
        return {"sku": self.sku, "campo": self.campo, "despues": self.despues,
                "categoria_id": self.categoria_id}


@dataclass
class Revision:
    cambios: list = field(default_factory=list)
    errores: list = field(default_factory=list)  # (fila, sku, mensaje)
    sin_cambio: int = 0


def proponer(producto, *, mascota="", categoria=None, descripcion="") -> list[Cambio]:
    """Lista de cambios para un producto. Vacío = no tocar (la regla)."""
    cambios = []
    if mascota and mascota != producto.mascota:
        cambios.append(Cambio(producto.sku, producto.nombre, "mascota", producto.mascota, mascota, producto=producto))
    if categoria is not None and categoria.id != producto.categoria_id:
        cambios.append(Cambio(producto.sku, producto.nombre, "categoria",
                              etiqueta_categoria(producto.categoria), etiqueta_categoria(categoria),
                              categoria_id=categoria.id, producto=producto))
    descripcion = (descripcion or "").strip()
    if descripcion and descripcion != (producto.descripcion or "").strip():
        cambios.append(Cambio(producto.sku, producto.nombre, "descripcion", producto.descripcion or "", descripcion, producto=producto))
    return cambios


@transaction.atomic
def aplicar(empresa, cambios: list[dict]) -> int:
    """Aplica cambios ya revisados. Devuelve cuántos productos se tocaron.

    Se vuelve a validar todo acá: entre la vista previa y la confirmación
    puede haber pasado tiempo, y lo que llega del navegador no se asume
    bueno solo porque lo firmamos nosotros."""
    por_sku = {}
    for c in cambios:
        por_sku.setdefault(c["sku"], []).append(c)
    tocados = 0
    for sku, lista in por_sku.items():
        producto = Producto.objects.select_for_update().filter(empresa=empresa, sku=sku).first()
        if producto is None:
            continue
        campos = []
        for c in lista:
            if c["campo"] == "mascota" and c["despues"] in MASCOTAS:
                producto.mascota = c["despues"]
                campos.append("mascota")
            elif c["campo"] == "categoria" and c.get("categoria_id"):
                cat = Categoria.objects.filter(pk=c["categoria_id"]).first()
                if cat is not None:
                    producto.categoria = cat
                    campos.append("categoria")
            elif c["campo"] == "descripcion" and (c["despues"] or "").strip():
                producto.descripcion = c["despues"].strip()
                campos.append("descripcion")
        if campos:
            producto.save(update_fields=campos + ["actualizado_en"])
            tocados += 1
    return tocados


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------

_COLUMNAS = ["Código (SKU)", "Nombre", "Existencia", "Mascota", "Categoría", "Descripción", "Le falta"]


def exportar_excel(productos) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = HOJA_PRODUCTOS
    ws.append(_COLUMNAS)
    negrita = Font(bold=True, color="FFFFFF")
    for celda in ws[1]:
        celda.font = negrita
        celda.fill = PatternFill("solid", fgColor="0B3161")
    editable = PatternFill("solid", fgColor="FFF7E6")
    n = 1
    for p in productos:
        n += 1
        ws.append([
            p.sku, p.nombre, float(p.stock_actual), p.mascota,
            etiqueta_categoria(p.categoria), p.descripcion or "",
            ", ".join(ETIQUETA_CAMPO[f] for f in faltantes(p)),
        ])
        ws.cell(row=n, column=1).number_format = "@"
        for col in (4, 5, 6):
            ws.cell(row=n, column=col).fill = editable
        ws.cell(row=n, column=6).alignment = Alignment(wrap_text=True, vertical="top")
    for letra, ancho in zip("ABCDEFG", (14, 46, 11, 16, 38, 70, 30)):
        ws.column_dimensions[letra].width = ancho
    ws.freeze_panes = "C2"

    op = wb.create_sheet(HOJA_OPCIONES)
    op.append(["Mascota", "Categoría"])
    for celda in op[1]:
        celda.font = Font(bold=True)
    cats = [etq for _, etq in opciones_categoria()]
    for i in range(max(len(MASCOTAS), len(cats))):
        op.append([MASCOTAS[i] if i < len(MASCOTAS) else None, cats[i] if i < len(cats) else None])
    op.column_dimensions["A"].width = 16
    op.column_dimensions["B"].width = 44

    if n > 1:
        dv_m = DataValidation(type="list", formula1=f"={HOJA_OPCIONES}!$A$2:$A${len(MASCOTAS) + 1}", allow_blank=True)
        dv_c = DataValidation(type="list", formula1=f"={HOJA_OPCIONES}!$B$2:$B${len(cats) + 1}", allow_blank=True)
        ws.add_data_validation(dv_m)
        ws.add_data_validation(dv_c)
        dv_m.add(f"D2:D{n}")
        dv_c.add(f"E2:E{n}")

    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


def _texto_celda(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def revisar_excel(empresa, archivo) -> Revision:
    """Lee el Excel subido y arma la lista de cambios, SIN aplicar nada."""
    from openpyxl import load_workbook

    rev = Revision()
    try:
        wb = load_workbook(archivo, read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 — cualquier archivo que no sea xlsx válido
        rev.errores.append((None, "", "El archivo no se pudo abrir como Excel (.xlsx)."))
        return rev
    ws = wb[HOJA_PRODUCTOS] if HOJA_PRODUCTOS in wb.sheetnames else wb.worksheets[0]
    filas = ws.iter_rows(values_only=True)
    encabezado = [_plano(v) for v in (next(filas, None) or [])]

    def col(*nombres):
        for nombre in nombres:
            if _plano(nombre) in encabezado:
                return encabezado.index(_plano(nombre))
        return None

    i_sku = col("Código (SKU)", "SKU", "Código", "Codigo")
    i_m, i_c, i_d = col("Mascota"), col("Categoría", "Categoria"), col("Descripción", "Descripcion")
    if i_sku is None:
        rev.errores.append((1, "", "No encontré la columna del código (SKU). Usá el Excel que descarga esta pantalla."))
        return rev

    productos = {p.sku: p for p in Producto.objects.filter(empresa=empresa).select_related("categoria", "categoria__padre")}
    vistos = set()
    for n, fila in enumerate(filas, start=2):
        def celda(i):
            return _texto_celda(fila[i]) if i is not None and i < len(fila) else ""

        sku = celda(i_sku)
        if not sku:
            continue
        if sku in vistos:
            rev.errores.append((n, sku, "El código aparece dos veces en el Excel; se usó la primera fila."))
            continue
        vistos.add(sku)
        producto = productos.get(sku)
        if producto is None:
            rev.errores.append((n, sku, "Ese código no existe en el ERP."))
            continue

        mascota = ""
        if celda(i_m):
            mascota = normalizar_mascota(celda(i_m)) or ""
            if not mascota:
                rev.errores.append((n, sku, f"Mascota «{celda(i_m)}» no válida. Opciones: {', '.join(MASCOTAS)}."))
        categoria = None
        if celda(i_c):
            categoria = buscar_categoria(celda(i_c))
            if categoria is None:
                rev.errores.append((n, sku, f"Categoría «{celda(i_c)}» no existe. Copiala de la hoja «{HOJA_OPCIONES}»."))

        cambios = proponer(producto, mascota=mascota, categoria=categoria, descripcion=celda(i_d))
        if cambios:
            rev.cambios.extend(cambios)
        else:
            rev.sin_cambio += 1
    return rev
