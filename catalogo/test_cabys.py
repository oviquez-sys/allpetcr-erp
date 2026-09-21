"""CABYS por categoría (21/09/2026) — catalogo/cabys.py y asignar_cabys."""
import tempfile
from decimal import Decimal
from io import StringIO
from pathlib import Path

from django.core.management import call_command
from django.test import TestCase

from core.models import AuditLog, Empresa

from . import cabys
from .models import Categoria, Impuesto, Producto


class TablaDeCodigos(TestCase):
    """La tabla es la única fuente de códigos: si una regla apunta a un
    código que no está en CODIGOS (verificado en Hacienda), es un código
    inventado."""

    def test_todos_los_codigos_son_de_13_digitos(self):
        for codigo in cabys.CODIGOS:
            self.assertRegex(codigo, r"^\d{13}$")

    def test_ninguna_regla_usa_un_codigo_fuera_de_la_tabla(self):
        usados = {c for c, _, _ in cabys._POR_RAIZ.values()}
        usados |= {c for c, _, _ in cabys._POR_SUBCATEGORIA.values()}
        usados |= {r[1] for reglas in cabys._REGLAS.values() for r in reglas}
        self.assertEqual(usados - set(cabys.CODIGOS), set())


class _Base(TestCase):
    def setUp(self):
        self.empresa = Empresa.objects.create(nombre="ALLPETCR.COM")
        self.iva = Impuesto.objects.create(nombre="IVA general", tarifa=Decimal("13.00"))
        self._n = 0

    def cat(self, raiz, hija=None):
        padre, _ = Categoria.objects.get_or_create(nombre=raiz)
        if hija is None:
            return padre
        return Categoria.objects.get_or_create(nombre=hija, defaults={"padre": padre})[0]

    def prod(self, nombre, categoria, **extra):
        self._n += 1
        datos = dict(empresa=self.empresa, sku=f"T{self._n:03d}", nombre=nombre, categoria=categoria,
                     precio_venta=Decimal("1000"), impuesto=self.iva)
        datos.update(extra)
        return Producto.objects.create(**datos)


class Propuestas(_Base):
    def _codigo(self, nombre, categoria):
        return cabys.proponer(self.prod(nombre, categoria)).codigo

    def test_champu_y_cepillo_de_la_misma_categoria_van_distinto(self):
        higiene = self.cat("Higiene y aseo", "Baño")
        self.assertEqual(self._codigo("Champú para perro 500 ml", higiene), "3532307000200")
        self.assertEqual(self._codigo("CEPILLO de doble cara", higiene), "3899302029900")
        self.assertEqual(self._codigo("Toallitas húmedas", higiene), "3532307009900")

    def test_correa_y_collar_antipulgas(self):
        paseo = self.cat("Paseo", "Collares")
        self.assertEqual(self._codigo("Correa retráctil 5 m", paseo), "2921001000000")
        self.assertEqual(self._codigo("Collar antipulgas gato", paseo), "3466100990400")

    def test_alimento_humedo_por_subcategoria(self):
        self.assertEqual(self._codigo("Royal Canin adulto", self.cat("Alimento", "Alimento húmedo")), "2331100000100")
        self.assertEqual(self._codigo("Royal Canin adulto 15 kg", self.cat("Alimento", "Alimento seco")), "2331100000200")

    def test_alimento_de_peces_en_acuario(self):
        acuario = self.cat("Acuario")
        self.assertEqual(self._codigo("Alimento en hojuelas 50 g", acuario), "2331901000100")
        self.assertEqual(cabys.proponer(self.prod("Planta artificial", acuario)).confianza, cabys.BAJA)

    def test_arenero_colgado_de_la_raiz_no_se_vuelve_arena(self):
        raiz = self.cat("Arena y sanitarios")
        self.assertEqual(self._codigo("Arenero cerrado con tapa", raiz), "3694000999900")
        self.assertEqual(self._codigo("Arena aglomerante 10 kg", raiz), "1540002009900")

    def test_sin_categoria_no_se_inventa_nada(self):
        self.assertIsNone(cabys.proponer(self.prod("Lámpara solar", None)))


class ComandoAsignar(_Base):
    def setUp(self):
        super().setUp()
        self.dir = tempfile.TemporaryDirectory()
        self.excel = Path(self.dir.name) / "cabys.xlsx"
        self.juguete = self.prod("Pelota con sonido", self.cat("Juguetes", "Pelotas"))
        self.correa = self.prod("Correa 1,2 m", self.cat("Paseo", "Correas"))
        self.huerfano = self.prod("Botón", None)

    def tearDown(self):
        self.dir.cleanup()

    def correr(self, *args):
        call_command("asignar_cabys", "--excel", str(self.excel), *args, stdout=StringIO())

    def test_asigna_y_deja_excel_para_el_contador(self):
        self.correr()
        self.juguete.refresh_from_db()
        self.correa.refresh_from_db()
        self.huerfano.refresh_from_db()
        self.assertEqual(self.juguete.cabys, "3856000009900")
        self.assertEqual(self.correa.cabys, "2921001000000")
        self.assertEqual(self.huerfano.cabys, "")

        from openpyxl import load_workbook
        hoja = load_workbook(self.excel).worksheets[0]
        filas = list(hoja.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(filas), 3)
        # Lo que el contador tiene que mirar primero va arriba: el sin código.
        self.assertEqual(filas[0][0], self.huerfano.sku)
        self.assertEqual(hoja.cell(2, 6).number_format, "@")

    def test_simulacion_no_escribe(self):
        self.correr("--dry-run")
        self.juguete.refresh_from_db()
        self.assertEqual(self.juguete.cabys, "")
        self.assertTrue(self.excel.exists())

    def test_no_pisa_un_codigo_puesto_a_mano(self):
        Producto.objects.filter(pk=self.juguete.pk).update(cabys="3852002000000")
        self.correr()
        self.juguete.refresh_from_db()
        self.assertEqual(self.juguete.cabys, "3852002000000")
        self.correr("--reemplazar")
        self.juguete.refresh_from_db()
        self.assertEqual(self.juguete.cabys, "3856000009900")

    def test_producto_con_otra_tarifa_no_recibe_codigo(self):
        reducida = Impuesto.objects.create(nombre="Reducida", tarifa=Decimal("1.00"))
        Producto.objects.filter(pk=self.correa.pk).update(impuesto=reducida)
        self.correr()
        self.correa.refresh_from_db()
        self.assertEqual(self.correa.cabys, "")

    def test_queda_en_la_bitacora_de_auditoria(self):
        antes = AuditLog.objects.filter(tabla="catalogo.producto").count()
        self.correr()
        self.assertGreater(AuditLog.objects.filter(tabla="catalogo.producto").count(), antes)

    def test_correr_dos_veces_no_cambia_nada_la_segunda(self):
        self.correr()
        antes = AuditLog.objects.filter(tabla="catalogo.producto").count()
        self.correr()
        self.assertEqual(AuditLog.objects.filter(tabla="catalogo.producto").count(), antes)
