"""Aplica la lista de precios aprobada por Oscar y pasa la empresa a régimen
tradicional, todo junto (20/09/2026).

Contexto: AllPetCR resultó ser régimen tradicional desde el primer día y los
precios se habían puesto sin IVA. Oscar aprobó la "versión Medio" del plan de
precios (ver claude/regimen-tradicional-impacto.md en el proyecto): cada
precio ya trae el IVA incluido, subido según el precio del chino y terminado
en precio bonito (catalogo/precios_bonitos.py).

Por qué las dos cosas en UNA transacción: si cambia el régimen y no los
precios, cada venta le regala el IVA a Hacienda; si cambian los precios y no
el régimen, el sistema no separa el IVA y los números del mes mienten.

    python manage.py aplicar_precios_iva --excel data/PRECIOS_NUEVOS_IVA.xlsx --dry-run
    python manage.py aplicar_precios_iva --excel data/PRECIOS_NUEVOS_IVA.xlsx

El Excel es el que aprobó Oscar: columnas "Código", "Precio hoy" y
"PRECIO NUEVO (con IVA)" en la primera hoja.

Seguridad:
- Nunca BAJA un precio: si el precio del sistema ya es igual o mayor al
  nuevo (alguien lo cambió después del plan), lo deja y lo reporta.
- Un código repetido en el Excel (misma mercadería comprada a dos costos)
  toma el precio más alto: el sistema tiene un solo precio por producto.
- Se niega a correr dos veces: cada cambio queda en el historial de precios
  con MOTIVO, y eso es la marca.
- Se niega si la migración catalogo.0007 no está en esa base (los cambios
  del ERP todavía no se subieron): con el código viejo, un producto sin
  tarifa se vendería sin IVA.
"""
from decimal import Decimal

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.db.migrations.recorder import MigrationRecorder

from catalogo.models import CambioPrecio, Producto
from catalogo.services import cambiar_precio
from core.models import Empresa

MOTIVO = "Precios nuevos con IVA 13 % incluido (régimen tradicional)"
COL_CODIGO = "Código"
COL_HOY = "Precio hoy"
COL_NUEVO = "PRECIO NUEVO (con IVA)"


def leer_excel(ruta):
    """{sku: (precio_nuevo, precio_hoy)} de la primera hoja del Excel aprobado."""
    from openpyxl import load_workbook

    hoja = load_workbook(ruta, read_only=True, data_only=True).worksheets[0]
    filas = hoja.iter_rows(values_only=True)
    cabecera = [str(c).strip() if c is not None else "" for c in next(filas)]
    faltan = [c for c in (COL_CODIGO, COL_HOY, COL_NUEVO) if c not in cabecera]
    if faltan:
        raise CommandError(f"Al Excel le faltan las columnas: {', '.join(faltan)}.")
    i_cod, i_hoy, i_nuevo = (cabecera.index(c) for c in (COL_CODIGO, COL_HOY, COL_NUEVO))
    precios = {}
    for fila in filas:
        if not fila or fila[i_cod] in (None, ""):
            continue
        sku = str(fila[i_cod]).strip()
        nuevo = Decimal(str(fila[i_nuevo])).quantize(Decimal("1"))
        hoy = Decimal(str(fila[i_hoy] or 0)).quantize(Decimal("1"))
        if nuevo <= 0:
            raise CommandError(f"El código {sku} trae un precio nuevo inválido: {fila[i_nuevo]}.")
        anterior = precios.get(sku)
        if anterior is None or nuevo > anterior[0]:
            precios[sku] = (nuevo, hoy)
    if not precios:
        raise CommandError("El Excel no trae ningún precio.")
    return precios


class Command(BaseCommand):
    help = "Aplica los precios nuevos con IVA del Excel aprobado y pasa la empresa a régimen tradicional."

    def add_arguments(self, parser):
        parser.add_argument("--excel", required=True, help="Excel aprobado con los precios nuevos.")
        parser.add_argument("--dry-run", action="store_true", help="Solo muestra; no escribe nada.")

    def handle(self, *args, **opciones):
        empresa = Empresa.objects.first()
        if empresa is None:
            raise CommandError("No hay empresa configurada.")
        if not MigrationRecorder.Migration.objects.filter(
            app="catalogo", name="0007_iva_general_en_productos"
        ).exists():
            raise CommandError(
                "Esta base todavía no tiene los cambios del 20/09/2026. Primero subí los cambios "
                "del ERP (SUBIR_CAMBIOS.bat), esperá que termine el despliegue y volvé a correr esto."
            )
        if CambioPrecio.objects.filter(producto__empresa=empresa, motivo=MOTIVO).exists():
            raise CommandError("Los precios nuevos YA se aplicaron en esta base. No se hizo nada.")

        precios = leer_excel(opciones["excel"])
        productos = {p.sku: p for p in Producto.objects.filter(empresa=empresa, sku__in=precios)}

        a_cambiar, no_bajan, no_existen = [], [], []
        for sku, (nuevo, hoy_plan) in sorted(precios.items()):
            p = productos.get(sku)
            if p is None:
                no_existen.append(sku)
            elif nuevo <= p.precio_venta:
                no_bajan.append((p, nuevo))
            else:
                a_cambiar.append((p, nuevo, hoy_plan))

        w = self.stdout.write
        w(f"\n  Empresa: {empresa.nombre} — régimen actual: {empresa.get_regimen_display()}")
        w(f"  Precios en el Excel: {len(precios)}")
        w(f"  Van a cambiar:       {len(a_cambiar)}")
        if a_cambiar:
            antes = sum(p.precio_venta for p, _, _ in a_cambiar)
            despues = sum(n for _, n, _ in a_cambiar)
            w(f"  Suma de esos precios: ₡{_m(antes)} → ₡{_m(despues)} ({(despues / antes - 1) * 100:+.0f} %)")
            w("\n  Algunos ejemplos:")
            for p, nuevo, _ in a_cambiar[:: max(1, len(a_cambiar) // 12)][:12]:
                w(f"    {p.sku:>8}  {p.nombre[:34]:34}  ₡{_m(p.precio_venta):>7} → ₡{_m(nuevo):>7}")
        distintos = [(p, h) for p, _, h in a_cambiar if h and h != p.precio_venta]
        if distintos:
            w(self.style.WARNING(f"\n  OJO: {len(distintos)} producto(s) tienen hoy otro precio que el del plan (igual suben):"))
            for p, h in distintos[:15]:
                w(f"    {p.sku}  {p.nombre[:40]}  plan decía ₡{_m(h)}, sistema tiene ₡{_m(p.precio_venta)}")
        if no_bajan:
            w(self.style.WARNING(f"\n  No se tocan {len(no_bajan)} (su precio actual ya es igual o mayor):"))
            for p, n in no_bajan[:15]:
                w(f"    {p.sku}  {p.nombre[:40]}  sistema ₡{_m(p.precio_venta)}, plan ₡{_m(n)}")
        if no_existen:
            w(self.style.WARNING(f"\n  No están en el sistema {len(no_existen)} código(s): {', '.join(no_existen[:20])}"))
        w("\n  Además: la empresa pasa a RÉGIMEN TRADICIONAL.")

        if opciones["dry_run"]:
            w(self.style.WARNING("\n  SIMULACIÓN: no se escribió nada."))
            return

        with transaction.atomic():
            for p, nuevo, _ in a_cambiar:
                cambiar_precio(producto=p, nuevo_precio=nuevo, motivo=MOTIVO)
            empresa.regimen = Empresa.Regimen.TRADICIONAL
            empresa.save(update_fields=["regimen"])
        w(self.style.SUCCESS(
            f"\n  LISTO: {len(a_cambiar)} precios actualizados y la empresa quedó en régimen tradicional."
        ))


def _m(valor):
    return f"{valor:,.0f}".replace(",", ".")
