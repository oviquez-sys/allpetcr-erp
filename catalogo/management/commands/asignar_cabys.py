"""Carga el código CABYS de cada producto según su categoría (21/09/2026).

    python manage.py asignar_cabys --dry-run    # ver qué haría, sin escribir
    python manage.py asignar_cabys              # aplicarlo

Los códigos y las reglas están en catalogo/cabys.py (verificados en la API
de Hacienda). Siempre deja un Excel para el contador —por defecto
data/CABYS_PARA_EL_CONTADOR.xlsx— con cada producto, su código, la
descripción oficial y qué tan seguro es; los de confianza baja arriba.

Seguridad:
- Un producto que YA tiene CABYS no se toca (lo puso alguien a mano, o lo
  corrigió el contador). `--reemplazar` los pisa, y solo tiene sentido si se
  cambiaron las reglas.
- Un producto cuya tarifa no es 13 % no recibe código: todos los códigos de
  la tabla son de 13 %, y un código con otra tarifa que la del producto haría
  que Hacienda rechace la factura. Se reporta para que lo vea el contador.
- Todo en una transacción: o quedan todos, o ninguno.
- Cada cambio queda en la bitácora de auditoría (señal de catalogo.producto).
"""
from decimal import Decimal
from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand
from django.db import transaction

from catalogo import cabys
from catalogo.models import Producto, validar_cabys

TARIFA_DE_LA_TABLA = Decimal("13.00")
_ORDEN_CONFIANZA = {cabys.BAJA: 0, cabys.MEDIA: 1, cabys.ALTA: 2}


class Command(BaseCommand):
    help = "Asigna el código CABYS de cada producto por su categoría."

    def add_arguments(self, parser):
        parser.add_argument("--dry-run", action="store_true", help="No escribe nada en la base.")
        parser.add_argument("--reemplazar", action="store_true",
                            help="También cambia los productos que ya tienen CABYS.")
        parser.add_argument("--excel", default=str(Path(settings.BASE_DIR) / "data" / "CABYS_PARA_EL_CONTADOR.xlsx"),
                            help="Dónde dejar el Excel para el contador.")

    def handle(self, *args, **op):
        filas = []
        asignados = ya_tenian = sin_regla = otra_tarifa = 0

        with transaction.atomic():
            productos = (Producto.objects.filter(activo=True)
                         .select_related("categoria__padre", "impuesto").order_by("sku"))
            for p in productos:
                propuesta = cabys.proponer(p)
                fila = {
                    "sku": p.sku, "nombre": p.nombre,
                    "categoria": self._etiqueta(p.categoria),
                    "existencia": p.stock_actual,
                    "tarifa": p.tarifa_iva,
                    "anterior": p.cabys,
                }
                if propuesta is None:
                    sin_regla += 1
                    filas.append({**fila, "cabys": p.cabys, "descripcion": "",
                                  "confianza": "", "estado": "SIN CÓDIGO: falta categoría"})
                    continue
                fila.update(cabys=propuesta.codigo, descripcion=propuesta.descripcion,
                            confianza=propuesta.confianza, motivo=propuesta.motivo)
                if p.tarifa_iva != TARIFA_DE_LA_TABLA:
                    otra_tarifa += 1
                    fila.update(cabys=p.cabys, estado=f"NO SE ASIGNÓ: el producto tiene IVA {p.tarifa_iva} %")
                elif p.cabys and not op["reemplazar"]:
                    ya_tenian += 1
                    fila.update(cabys=p.cabys, descripcion=cabys.CODIGOS.get(p.cabys, "(puesto a mano)"),
                                estado="Ya tenía código: no se tocó")
                else:
                    asignados += 1
                    fila["estado"] = "Asignado" if not p.cabys else "Reemplazado"
                    if p.cabys != propuesta.codigo:
                        validar_cabys(propuesta.codigo)  # 13 dígitos, o revienta antes de escribir
                        p.cabys = propuesta.codigo
                        p.save(update_fields=["cabys", "actualizado_en"])
                filas.append(fila)

            if op["dry_run"]:
                transaction.set_rollback(True)

        ruta = self._excel(filas, Path(op["excel"]))
        w = self.stdout.write
        w(self.style.SUCCESS("\nSIMULACION (no se escribio nada)" if op["dry_run"] else "\nCABYS cargados"))
        w(f"  Productos activos ............ {len(filas)}")
        w(f"  Con codigo asignado .......... {asignados}")
        w(f"  Ya tenian codigo (no tocados)  {ya_tenian}")
        w(f"  Sin codigo: falta categoria .. {sin_regla}")
        w(f"  Sin codigo: IVA distinto 13 %  {otra_tarifa}")
        conteo = {}
        for f in filas:
            if f.get("confianza"):
                conteo[f["confianza"]] = conteo.get(f["confianza"], 0) + 1
        w(f"  Confianza alta / media / baja  {conteo.get('alta', 0)} / {conteo.get('media', 0)} / {conteo.get('baja', 0)}")
        w(f"\n  Excel para el contador: {ruta}")

    @staticmethod
    def _etiqueta(cat):
        if cat is None:
            return ""
        return f"{cat.padre.nombre} > {cat.nombre}" if cat.padre_id else cat.nombre

    def _excel(self, filas, ruta):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        ruta.parent.mkdir(parents=True, exist_ok=True)
        libro = Workbook()
        hoja = libro.active
        hoja.title = "CABYS"
        cabecera = ["Código (SKU)", "Producto", "Categoría", "Existencia", "IVA %",
                    "CABYS", "Descripción oficial de Hacienda", "Confianza", "Por qué", "Estado",
                    "CABYS corregido por el contador"]
        hoja.append(cabecera)
        for c in hoja[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="0B3161")
        colores = {cabys.BAJA: "FDE2E1", cabys.MEDIA: "FFF4D6"}
        filas = sorted(filas, key=lambda f: (_ORDEN_CONFIANZA.get(f.get("confianza"), -1), f["categoria"], f["nombre"]))
        for f in filas:
            hoja.append([f["sku"], f["nombre"], f["categoria"], float(f["existencia"]), float(f["tarifa"]),
                         f.get("cabys") or "", f.get("descripcion", ""), f.get("confianza", ""),
                         f.get("motivo", ""), f.get("estado", ""), ""])
            color = colores.get(f.get("confianza")) or ("FDE2E1" if not f.get("cabys") else None)
            if color:
                for c in hoja[hoja.max_row]:
                    c.fill = PatternFill("solid", fgColor=color)
            hoja.cell(hoja.max_row, 6).number_format = "@"  # que Excel no lo vuelva 2,33E+12
        for i, ancho in enumerate([12, 45, 32, 10, 7, 16, 60, 10, 30, 32, 22], start=1):
            hoja.column_dimensions[get_column_letter(i)].width = ancho
        hoja.freeze_panes = "A2"
        hoja.auto_filter.ref = hoja.dimensions

        nota = libro.create_sheet("Cómo revisar")
        for linea in [
            "Códigos CABYS asignados por el ERP de AllPetCR según la categoría de cada producto.",
            "Cada código se verificó el 21/09/2026 en la API oficial de Hacienda. Todos llevan IVA 13 %.",
            "Rojo = confianza baja o sin código; amarillo = confianza media; blanco = alta.",
            "Si un código no corresponde, escriba el correcto en la última columna y se corrige en el ERP.",
        ]:
            nota.append([linea])
        nota.column_dimensions["A"].width = 110
        libro.save(ruta)
        return ruta
