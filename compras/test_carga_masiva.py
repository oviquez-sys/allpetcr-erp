"""INV-01 (auditoría 26/09/2026): carga masiva de mercadería desde Excel/CSV.

Lo que fijan estas pruebas: subir NO guarda nada; con un error no se puede
confirmar; al confirmar entra todo junto (productos nuevos, compra, stock,
costo, asiento y precios) y lo que se confirma es lo que se mostró.
"""
import io
from decimal import Decimal
from unittest import mock

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from catalogo.codigos import es_interno
from catalogo.models import CambioPrecio, Categoria, Producto
from impresion.models import TrabajoImpresion

from .models import Compra
from .tests import BaseCompras


def _excel(filas, encabezado=("Código de barras", "Nombre", "Cantidad", "Costo unitario", "Precio de venta",
                              "Bonificadas", "Marca", "Categoría", "Presentación", "Mascota")):
    wb = Workbook()
    ws = wb.active
    ws.append(list(encabezado))
    for f in filas:
        ws.append(list(f))
    salida = io.BytesIO()
    wb.save(salida)
    return SimpleUploadedFile("mercaderia.xlsx", salida.getvalue())


class CargaMasiva(BaseCompras):
    def setUp(self):
        super().setUp()
        self.producto.codigo_barras = "7501234567895"
        self.producto.save()
        Categoria.objects.create(nombre="Juguetes")
        self.client.force_login(self.usuario)

    def subir(self, archivo, **extra):
        datos = {"proveedor_id": self.proveedor.pk, "factura_proveedor": "F-900", "forma_pago": "CON",
                 "archivo": archivo, **extra}
        return self.client.post(reverse("compras:carga_masiva"), datos)

    def test_la_plantilla_se_descarga_y_se_abre(self):
        r = self.client.get(reverse("compras:carga_masiva_plantilla"))
        wb = load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb.active["A1"].value, "Código de barras")

    def test_subir_no_guarda_nada(self):
        r = self.subir(_excel([["7501234567895", "", 5, 9000, "", 0],
                               ["", "Pelota nueva", 10, 500, 1500, 2, "", "Juguetes", "", "perro"]]))
        self.assertContains(r, "Todo en orden")
        self.assertEqual(Compra.objects.count(), 0)
        self.assertFalse(Producto.objects.filter(nombre="Pelota nueva").exists())

    def test_con_errores_no_se_puede_confirmar(self):
        r = self.subir(_excel([["", "Sin precio", 3, 100, "", 0],
                               ["7501234567895", "", "tres", 100, "", 0],
                               ["", "Otra", 1, 1, 100, 0, "", "No existe"]]))
        self.assertContains(r, "falta el precio de venta")
        self.assertContains(r, "no es un número")
        self.assertContains(r, "no existe en el ERP")
        self.assertNotIn("firma", r.context)

    def confirmar(self, archivo):
        r = self.subir(archivo)
        return self.client.post(reverse("compras:carga_masiva_confirmar"), {"firma": r.context["firma"]})

    def test_confirmar_carga_todo(self):
        r = self.confirmar(_excel([
            ["7501234567895", "", 5, 9000, 19500, 1],
            ["", "Pelota nueva", 10, 500, 1500, 2, "Kong", "Juguetes", "Talla S", "Perro"],
        ]))
        self.assertEqual(r.status_code, 200)
        compra = Compra.objects.get()
        self.assertEqual(compra.estado, Compra.Estado.RECIBIDA)
        self.assertEqual(compra.factura_proveedor, "F-900")
        self.assertEqual(compra.total, Decimal("50000"))            # 5×9000 + 10×500
        self.producto.refresh_from_db()
        self.assertEqual(self.producto.stock_actual, Decimal("16"))   # 10 + 5 + 1 bonificada
        self.assertEqual(self.producto.precio_venta, Decimal("19500"))
        self.assertTrue(CambioPrecio.objects.filter(producto=self.producto).exists())
        nuevo = Producto.objects.get(nombre="Pelota nueva")
        self.assertEqual(nuevo.stock_actual, Decimal("12"))
        self.assertTrue(es_interno(nuevo.codigo_barras))
        self.assertEqual((nuevo.marca, nuevo.mascota, nuevo.categoria.nombre), ("Kong", "Perro", "Juguetes"))

    def test_la_misma_factura_no_entra_dos_veces(self):
        self.confirmar(_excel([["7501234567895", "", 5, 9000, "", 0]]))
        self.confirmar(_excel([["7501234567895", "", 5, 9000, "", 0]]))
        self.assertEqual(Compra.objects.count(), 1)

    def test_si_algo_falla_al_confirmar_no_queda_nada(self):
        r = self.subir(_excel([["", "Pelota nueva", 10, 500, 1500, 0]]))
        with mock.patch("compras.services.recibir_compra", side_effect=Exception("se cayó")):
            with self.assertRaises(Exception):
                self.client.post(reverse("compras:carga_masiva_confirmar"), {"firma": r.context["firma"]})
        self.assertFalse(Producto.objects.filter(nombre="Pelota nueva").exists())
        self.assertEqual(Compra.objects.count(), 0)

    def test_la_vista_previa_no_se_puede_alterar(self):
        r = self.client.post(reverse("compras:carga_masiva_confirmar"), {"firma": "inventada"})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(Compra.objects.count(), 0)

    def test_acepta_csv_con_punto_y_coma_y_montos_con_puntos(self):
        csv = "Código de barras;Nombre;Cantidad;Costo unitario;Precio de venta\n7501234567895;;2;2.500;\n"
        r = self.subir(SimpleUploadedFile("m.csv", csv.encode("utf-8")))
        self.assertContains(r, "Todo en orden")
        self.assertEqual(r.context["revision"].filas[0].costo_unitario, "2500.00")

    def test_nombre_ambiguo_es_error(self):
        Producto.objects.create(empresa=self.empresa, sku="X1", nombre="Alimento", precio_venta=1)
        r = self.subir(_excel([["", "Alimento", 1, 1, "", 0]]))
        self.assertContains(r, "productos que se llaman")


@override_settings(IMPRESION_FORZAR_AGENTE=True, IMPRESORA_ETIQUETAS="Etiquetas de prueba")
class EtiquetasDeLaCompra(BaseCompras):
    """INV-08: una etiqueta por unidad recibida, bonificadas incluidas."""

    def test_encola_una_por_unidad(self):
        from .services import crear_y_recibir_compra

        compra = crear_y_recibir_compra(
            proveedor=self.proveedor, sucursal=self.sucursal, usuario=self.usuario,
            lineas=[{"producto": self.producto, "cantidad": Decimal("3"), "cantidad_bonificada": Decimal("1"),
                     "costo_unitario": Decimal("1000")}])
        self.client.force_login(self.usuario)
        r = self.client.post(reverse("impresion:etiquetas_compra", args=[compra.pk]), HTTP_ACCEPT="application/json")
        self.assertTrue(r.json()["ok"])
        self.assertEqual(TrabajoImpresion.objects.filter(tipo=TrabajoImpresion.ETIQUETA).count(), 4)
