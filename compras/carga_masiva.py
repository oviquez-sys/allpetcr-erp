"""Carga masiva de mercadería desde Excel o CSV (auditoría 26/09/2026, INV-01).

POR QUÉ EXISTE
--------------
Hay mucha mercadería por ingresar y hasta hoy la única forma de cargar una
factura grande era un comando técnico (`cargar_compra_19_09`) que se corre
con un `.bat` y la contraseña de la base. Esta pantalla hace lo mismo desde
el navegador —también desde el celular— y con las mismas garantías:

1. **Nada se guarda al subir el archivo.** Primero se muestra cada fila:
   producto nuevo o existente, cantidad, costo, precio, y los errores.
2. **Con un solo error no se puede confirmar.** Se corrige el Excel y se
   vuelve a subir. Cargar "lo que sí sirve" deja la factura a medias y
   después nadie sabe qué faltó.
3. **Al confirmar entra todo junto o nada**: productos nuevos, la compra, el
   inventario, el costo promedio, el asiento y los precios nuevos, en una
   sola transacción.
4. **Lo que se confirma es exactamente lo que se vio**: la vista previa viaja
   firmada (igual que en «Completar catálogo»), así que no se puede alterar
   entre la revisión y la confirmación.

CÓMO SE RECONOCE UN PRODUCTO
----------------------------
Por el código interno (SKU) si la fila lo trae; si no, por el código de
barras (el de fábrica o el interno); si no, por el nombre exacto (sin
mayúsculas ni tildes). Si no aparece, es un producto NUEVO y entonces el
nombre y el precio de venta son obligatorios. Un nombre que coincide con dos
productos distintos es un error: adivinar cuál es ingresaría mercadería al
producto equivocado.

LAS COLUMNAS «CÓDIGO (SKU)» Y «DESCRIPCIÓN» (02/10/2026)
--------------------------------------------------------
Al cargar la primera factura con esta pantalla (Buen Amigo + Special Care)
faltaban dos cosas que las cargas anteriores por comando sí hacían:

- **Elegir el SKU del producto nuevo.** Sin la columna, el ERP le inventa
  uno («NP2610…») y se rompen dos convenciones: los productos de Buen Amigo
  usan como SKU el código REF del catálogo del proveedor (el que Oscar busca
  y el que sale en la factura), y las fotos se asocian POR SKU
  (`cargar_fotos_por_sku`): con un SKU inventado no hay forma de prepararle
  la foto antes de cargarlo.
- **La descripción pública.** Sin ella, cada producto nuevo salía sin texto
  en la web hasta llenarlo aparte en «Completar catálogo».

Las dos son opcionales: una plantilla vieja, sin esas columnas, funciona
igual que antes. En un producto que ya existe, la descripción solo se
completa si estaba vacía — misma regla que marca, presentación y mascota:
la carga nunca pisa un dato puesto a mano.

PAQUETE .ZIP: EXCEL + FOTOS, DESCUENTO, CABYS Y MARGEN MÍNIMO (27/10/2026)
-------------------------------------------------------------------------
Pedido de Oscar: "subo la factura y Claude se encarga de todo". Claude arma
un .zip por factura y Oscar lo sube acá. Hasta hoy, después de esta pantalla
quedaban tres pasos aparte, cada uno con su .bat y la contraseña de la base:
subir las fotos, poner el CABYS y revisar que ningún precio quedara sin
ganancia. Ahora:

- **Fotos en el mismo .zip.** Cada foto se llama como el SKU (o el código de
  barras) de su fila: `TOW-2573.webp`. En la vista previa se VE la foto al
  lado del producto, para notar una cambiada antes de confirmar. Una foto que
  no corresponde a ninguna fila es un error (seguro es un nombre mal escrito
  y el producto quedaría sin foto). A un producto que YA tiene foto no se le
  cambia: misma regla de "nunca pisa lo puesto a mano".
  Entre la vista previa y la confirmación las fotos esperan en el
  almacenamiento (`cargas_pendientes/<ficha>/`): son imágenes públicas, nunca
  el Excel (que lleva costos y viaja firmado, como antes).
- **Descuento de la factura**, el del pie, antes del IVA: se reparte entre
  las líneas igual que en "Recibir mercadería" (`services.crear_compra`).
- **Columna CABYS** (13 dígitos). Se pone al producto nuevo, o al existente
  que no tenía.
- **Margen mínimo** (`MARGEN_MINIMO_*`): la política de precios de Oscar es
  igualar el precio regular más bajo de la competencia, pero sin bajar del
  mínimo. Una fila por debajo NO bloquea: se pinta en amarillo para que él
  decida. El margen se calcula como en todo el ERP (`Producto.margen_pct`):
  sobre el precio SIN IVA y con el costo real (con el descuento y las
  unidades de regalo ya repartidos).
"""
import csv
import hashlib
import io
import posixpath
import unicodedata
import uuid
import zipfile
from dataclasses import asdict, dataclass, field
from datetime import timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation

from django.conf import settings
from django.core.exceptions import ValidationError
from django.core.files.base import ContentFile
from django.core.files.storage import default_storage
from django.db import transaction
from django.utils import timezone

from catalogo.completar import MASCOTAS, buscar_categoria, normalizar_mascota
from catalogo.models import TARIFA_GENERAL_IVA, Categoria, Producto
from catalogo.services import cambiar_precio
from core.imagenes import guardar_imagen_producto
from core.models import Empresa

from . import services

# Columnas de la plantilla, en orden. Las obligatorias van primero.
COLUMNAS = [
    ("codigo_barras", "Código de barras", "El de fábrica. Vacío si no tiene: el ERP le da uno interno."),
    ("nombre", "Nombre", "Obligatorio si el producto es nuevo."),
    ("cantidad", "Cantidad", "Unidades FACTURADAS. Obligatorio."),
    ("costo_unitario", "Costo unitario", "Costo de cada unidad, sin IVA si la factura lo trae aparte. Obligatorio."),
    ("precio_venta", "Precio de venta", "Con IVA incluido. Obligatorio si es nuevo; si es existente y cambia, se actualiza."),
    ("bonificadas", "Bonificadas", "Unidades de regalía del proveedor (el 1 del 12+1). Opcional."),
    ("marca", "Marca", "Opcional."),
    ("categoria", "Categoría", "Opcional. Como aparece en el ERP, ej. «Juguetes › Pelotas»."),
    ("presentacion", "Presentación", "Opcional, ej. «Talla M» o «15 kg»."),
    ("mascota", "Mascota", f"Opcional: {', '.join(MASCOTAS)}."),
    # Al final a propósito: así una plantilla descargada antes del 02/10/2026
    # sigue sirviendo tal cual (las columnas se reconocen por el título, no
    # por la posición).
    ("sku", "Código (SKU)", "Opcional. Código interno del producto, ej. el REF del catálogo del proveedor. "
                            "Si existe, la fila es de ese producto; si es nuevo, se crea con ese código."),
    ("descripcion", "Descripción", "Opcional. Texto público que se ve en la página web. "
                                   "En un producto existente solo se usa si no tenía descripción."),
    ("cabys", "CABYS", "Opcional. Código CABYS de Hacienda, 13 dígitos. "
                       "En un producto existente solo se usa si no tenía."),
]
MAX_FILAS = 1000

# Política de precios (decisión de Oscar, 27/10/2026): igualar el precio
# regular más bajo de la competencia, pero avisar si deja menos que esto.
# Alimento y snacks dejan poco por naturaleza; el resto, bastante más.
# Se pueden cambiar sin tocar código con las variables de entorno.
MARGEN_MINIMO_ALIMENTO = Decimal(str(getattr(settings, "CARGA_MARGEN_MINIMO_ALIMENTO", "15")))
MARGEN_MINIMO_OTROS = Decimal(str(getattr(settings, "CARGA_MARGEN_MINIMO_OTROS", "30")))
RAICES_ALIMENTO = ("Alimento", "Snacks y premios")

# Límites del .zip: lo que pesa una compra real (50 fotos WebP ≈ 5 MB) con
# holgura, y un tope a lo descomprimido para que un zip malicioso no llene
# la memoria del servidor.
EXTENSIONES_FOTO = (".jpg", ".jpeg", ".png", ".webp")
MAX_FOTOS = 600
MAX_DESCOMPRIMIDO = 300 * 1024 * 1024
CARPETA_PENDIENTES = "cargas_pendientes"


def _plano(texto) -> str:
    texto = " ".join(str(texto or "").split()).lower()
    return "".join(c for c in unicodedata.normalize("NFD", texto) if unicodedata.category(c) != "Mn")


# Encabezados que la gente escribe de verdad → campo.
_ALIAS = {_plano(etiqueta): clave for clave, etiqueta, _ in COLUMNAS}
_ALIAS.update({_plano(k): v for k, v in {
    "codigo": "codigo_barras", "código": "codigo_barras", "ean": "codigo_barras", "barras": "codigo_barras",
    "producto": "nombre", "descripcion": "nombre", "cant": "cantidad", "unidades": "cantidad",
    "costo": "costo_unitario", "precio": "precio_venta", "precio venta": "precio_venta",
    "bonificacion": "bonificadas", "regalo": "bonificadas", "regalia": "bonificadas", "especie": "mascota",
    "sku": "sku", "codigo interno": "sku", "codigo (ref)": "sku", "ref": "sku",
    "descripcion publica": "descripcion", "descripcion web": "descripcion",
}.items()})
# «Descripción» a secas siempre fue sinónimo de «Nombre» (mucha gente llama
# así a la columna del producto en la factura). Con la columna nueva hay que
# desempatar: si el archivo trae «Nombre», «Descripción» es la descripción
# pública; si no trae «Nombre», sigue siendo el nombre. Ver `revisar()`.
_ALIAS[_plano("Descripción")] = "descripcion"


@dataclass
class Fila:
    numero: int
    codigo_barras: str = ""
    nombre: str = ""
    cantidad: str = "0"
    costo_unitario: str = "0"
    precio_venta: str = ""
    bonificadas: str = "0"
    marca: str = ""
    categoria_id: int | None = None
    categoria_txt: str = ""
    presentacion: str = ""
    mascota: str = ""
    sku: str = ""
    descripcion: str = ""
    cabys: str = ""
    producto_id: int | None = None
    precio_actual: str = ""
    # Foto que llegó en el .zip: ruta en `cargas_pendientes/` y si se usa.
    foto_clave: str = ""      # nombre (sin extensión) de la foto en el zip
    foto: str = ""            # ruta de la foto en espera, si se va a usar
    foto_se_pone: bool = False
    # Margen con el precio que va a quedar (%, sin IVA) y el mínimo que aplica.
    margen: str = ""
    margen_minimo: str = ""
    errores: list = field(default_factory=list)
    avisos: list = field(default_factory=list)

    @property
    def nuevo(self):
        return self.producto_id is None

    @property
    def cambia_precio(self):
        return bool(self.precio_venta) and not self.nuevo and Decimal(self.precio_venta) != Decimal(self.precio_actual)

    @property
    def total(self):
        try:
            return Decimal(self.cantidad) * Decimal(self.costo_unitario)
        except InvalidOperation:
            return Decimal("0")


@dataclass
class Revision:
    filas: list = field(default_factory=list)
    errores_generales: list = field(default_factory=list)
    descuento: Decimal = Decimal("0")

    @property
    def con_errores(self):
        return [f for f in self.filas if f.errores]

    @property
    def con_avisos(self):
        return [f for f in self.filas if f.avisos]

    @property
    def fotos(self):
        return sum(1 for f in self.filas if f.foto_se_pone)

    @property
    def neto(self):
        return self.total - self.descuento

    @property
    def valida(self):
        return bool(self.filas) and not self.con_errores and not self.errores_generales

    @property
    def total(self):
        return sum((f.total for f in self.filas), Decimal("0"))

    @property
    def nuevos(self):
        return sum(1 for f in self.filas if f.nuevo)


def plantilla_excel() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Mercadería"
    ws.append([etiqueta for _, etiqueta, _ in COLUMNAS])
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="0B3161")
    ws.append(["7501234567895", "Ejemplo: Arnés acolchado talla M (borre esta fila)", 12, 2500, 6900, 1,
               "Marca X", "Paseo › Arneses", "Talla M", "Perro", "",
               "Arnés acolchado con cierre de clic y argolla metálica para la correa."])
    for i, ancho in enumerate([18, 42, 10, 14, 15, 12, 16, 26, 16, 14, 14, 60, 16], start=1):
        ws.column_dimensions[chr(64 + i)].width = ancho
    ayuda = wb.create_sheet("Instrucciones")
    ayuda.append(["Columna", "Qué va"])
    for _, etiqueta, texto in COLUMNAS:
        ayuda.append([etiqueta, texto])
    ayuda.column_dimensions["A"].width = 20
    ayuda.column_dimensions["B"].width = 90
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def _leer_filas(archivo):
    """Devuelve (encabezado, filas) de un .xlsx o un .csv."""
    nombre = (getattr(archivo, "name", "") or "").lower()
    if nombre.endswith(".csv"):
        texto = archivo.read().decode("utf-8-sig", errors="replace")
        dialecto = csv.Sniffer().sniff(texto[:2000], delimiters=",;\t") if texto.strip() else csv.excel
        filas = list(csv.reader(io.StringIO(texto), dialecto))
    else:
        from openpyxl import load_workbook

        wb = load_workbook(archivo, read_only=True, data_only=True)
        filas = [list(f) for f in wb.worksheets[0].iter_rows(values_only=True)]
    if not filas:
        return [], []
    return filas[0], filas[1:]


@dataclass
class Paquete:
    """Lo que trae lo subido: la hoja y, si fue un .zip, las fotos.

    `fotos`: nombre del archivo sin extensión, en minúsculas → (nombre
    original, bytes, extensión real con punto)."""
    hoja: object = None
    fotos: dict = field(default_factory=dict)
    errores: list = field(default_factory=list)


def es_zip(archivo) -> bool:
    return (getattr(archivo, "name", "") or "").lower().endswith(".zip")


def abrir_paquete(archivo) -> Paquete:
    """Separa la hoja y las fotos de un .zip. Un .xlsx o .csv suelto pasa tal cual.

    Las carpetas dentro del zip no importan (se usa solo el nombre del
    archivo): así sirve igual un zip armado a mano desde el Explorador."""
    if not es_zip(archivo):
        return Paquete(hoja=archivo)
    from .views import _validar_foto  # la misma validación de las otras pantallas de fotos

    paquete = Paquete()
    try:
        zf = zipfile.ZipFile(archivo)
        entradas = [i for i in zf.infolist() if not i.is_dir()]
    except zipfile.BadZipFile:
        paquete.errores.append("El .zip está dañado o no es un .zip.")
        return paquete
    if sum(i.file_size for i in entradas) > MAX_DESCOMPRIMIDO:
        paquete.errores.append("El .zip descomprimido pasa de 300 MB. Divídalo en partes.")
        return paquete
    hojas = []
    for info in entradas:
        nombre = posixpath.basename(info.filename.replace("\\", "/"))
        # Basura que agregan Mac y Windows al comprimir: no es parte de la compra.
        if not nombre or nombre.startswith((".", "~$")) or "__MACOSX" in info.filename:
            continue
        base, ext = posixpath.splitext(nombre)
        ext = ext.lower()
        if ext in (".xlsx", ".csv"):
            hojas.append((nombre, info))
        elif ext in EXTENSIONES_FOTO:
            if len(paquete.fotos) >= MAX_FOTOS:
                paquete.errores.append(f"El .zip trae más de {MAX_FOTOS} fotos. Divídalo en partes.")
                return paquete
            contenido = zf.read(info)
            try:
                ext_real = _validar_foto(contenido)
            except ValidationError as e:
                paquete.errores.append(f"Foto «{nombre}»: {' '.join(e.messages)}")
                continue
            clave = base.strip().lower()
            if clave in paquete.fotos:
                paquete.errores.append(f"Hay dos fotos para «{base}»: «{paquete.fotos[clave][0]}» y «{nombre}». "
                                       "Deje una sola.")
                continue
            paquete.fotos[clave] = (nombre, contenido, ext_real)
    if len(hojas) != 1:
        paquete.errores.append("El .zip tiene que traer UN solo Excel (.xlsx) o .csv con la mercadería"
                               + (f"; trae {len(hojas)}." if hojas else "; no trae ninguno."))
        return paquete
    nombre, info = hojas[0]
    hoja = io.BytesIO(zf.read(info))
    hoja.name = nombre
    paquete.hoja = hoja
    return paquete


def guardar_fotos_pendientes(rev: "Revision", paquete: Paquete) -> None:
    """Deja en espera las fotos que SÍ se van a usar (`foto_se_pone`).

    No toca ningún producto: es solo para que la confirmación (otra petición)
    las encuentre y para mostrarlas en la vista previa. Van con nombre al azar
    dentro de una carpeta al azar; las que no se usan ni se suben."""
    _limpiar_pendientes_viejos()
    # La fecha va adelante para poder limpiar las vistas previas abandonadas.
    carpeta = f"{CARPETA_PENDIENTES}/{timezone.now():%Y%m%d%H%M}-{uuid.uuid4().hex}"
    for f in rev.filas:
        if not f.foto_se_pone:
            continue
        _, contenido, ext = paquete.fotos[f.foto_clave]
        f.foto = default_storage.save(f"{carpeta}/{uuid.uuid4().hex[:12]}{ext}", ContentFile(contenido))


def _limpiar_pendientes_viejos():
    """Borra las fotos de vistas previas que nadie confirmó.

    La firma de la vista previa vence a la hora; con dos horas de margen,
    nada de lo que se borra puede estar esperando todavía una confirmación.
    Si falla no importa: son fotos públicas sueltas, no datos."""
    import logging

    limite = f"{timezone.now() - timedelta(hours=2):%Y%m%d%H%M}"
    try:
        carpetas, _ = default_storage.listdir(CARPETA_PENDIENTES)
    except Exception:  # noqa: BLE001 — la carpeta todavía no existe, o el bucket no responde
        logging.getLogger(__name__).info("Sin fotos en espera para limpiar.", exc_info=True)
        return
    for carpeta in carpetas:
        if carpeta[:12].isdigit() and carpeta[:12] < limite:
            try:
                _, archivos = default_storage.listdir(f"{CARPETA_PENDIENTES}/{carpeta}")
                for nombre in archivos:
                    default_storage.delete(f"{CARPETA_PENDIENTES}/{carpeta}/{nombre}")
            except Exception:  # noqa: BLE001 — se reintenta en la próxima vista previa
                logging.getLogger(__name__).warning("No se pudo limpiar %s", carpeta, exc_info=True)


def borrar_pendientes(filas):
    for datos in filas:
        ruta = datos.get("foto") if isinstance(datos, dict) else datos.foto
        if ruta:
            try:
                default_storage.delete(ruta)
            except Exception:  # noqa: BLE001 — una foto en espera que no se borra no daña nada
                import logging

                logging.getLogger(__name__).warning("No se pudo borrar la foto en espera %s", ruta, exc_info=True)


def _texto(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return " ".join(str(valor).split())


def _texto_con_renglones(valor) -> str:
    """Como `_texto`, pero respeta los saltos de línea: las descripciones
    públicas llevan una lista de viñetas («• Talla M», «• Ancho: 1 cm»), una
    por renglón, y el sitio las muestra así. Aplanarlas las juntaba en un
    solo párrafo ilegible."""
    if valor is None:
        return ""
    renglones = [" ".join(r.split()) for r in str(valor).replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    return "\n".join(r for r in renglones if r)


def _numero(valor):
    """Número de lo que escriba la gente: «2.500», «2500,50», «₡2 500». None si no es número."""
    texto = _texto(valor).replace("₡", "").replace("¢", "").replace(" ", "")
    if not texto:
        return Decimal("0")
    if "," in texto and "." in texto:          # 2.500,50 → 2500.50
        texto = texto.replace(".", "").replace(",", ".")
    elif "," in texto:                          # 2500,50 → 2500.50
        texto = texto.replace(",", ".")
    elif texto.count(".") == 1 and len(texto.split(".")[1]) == 3:
        texto = texto.replace(".", "")          # 2.500 (miles) → 2500
    try:
        return Decimal(texto).quantize(Decimal("0.01"))
    except InvalidOperation:
        return None


def revisar(empresa, archivo, *, fotos=None, descuento=Decimal("0")) -> Revision:
    """Lee el archivo y arma la vista previa. NO guarda nada.

    `fotos`: las del .zip (`Paquete.fotos`); `descuento`: el del pie de la
    factura, antes del IVA."""
    fotos = fotos or {}
    fotos_usadas = set()
    calculo = []  # (fila, tarifa de IVA, ¿es alimento?) para el margen, al final
    rev = Revision(descuento=Decimal(descuento or 0))
    try:
        encabezado, crudas = _leer_filas(archivo)
    except Exception:  # noqa: BLE001 — cualquier archivo ilegible es el mismo error para el usuario
        rev.errores_generales.append("No se pudo leer el archivo. Use la plantilla (.xlsx) o un .csv.")
        return rev
    columnas = {}
    titulos = {_plano(t) for t in encabezado}
    for i, titulo in enumerate(encabezado):
        clave = _ALIAS.get(_plano(titulo))
        # «Descripción» sin una columna «Nombre» al lado es el nombre del
        # producto, como siempre se aceptó (ver el comentario de _ALIAS).
        if clave == "descripcion" and _plano(titulo) == _plano("Descripción") and not (
                titulos & {_plano("Nombre"), _plano("Producto")}):
            clave = "nombre"
        if clave and clave not in columnas:
            columnas[clave] = i
    faltan = [et for c, et, _ in COLUMNAS[:4] if c not in columnas and c != "codigo_barras" and c != "nombre"]
    if not ({"codigo_barras", "nombre", "sku"} & set(columnas)):
        faltan.insert(0, "Código de barras o Nombre")
    if faltan:
        rev.errores_generales.append(f"Faltan columnas: {', '.join(faltan)}. Descargue la plantilla.")
        return rev

    productos = list(Producto.objects.filter(empresa=empresa, activo=True)
                     .select_related("categoria__padre", "impuesto"))
    por_codigo = {p.codigo_barras: p for p in productos if p.codigo_barras}
    por_sku = {p.sku: p for p in productos}
    # El SKU es único en TODA la base (no por empresa ni solo entre activos):
    # crear uno que ya usa un producto inactivo o de otra empresa revienta al
    # guardar. Se avisa en la vista previa, no a mitad de la carga.
    skus_ocupados = set(Producto.objects.values_list("sku", flat=True))
    codigos_ocupados = set(Producto.objects.exclude(codigo_barras="").values_list("codigo_barras", flat=True))
    skus_vistos = {}
    por_nombre = {}
    for p in productos:
        por_nombre.setdefault(_plano(p.nombre), []).append(p)
    codigos_vistos = {}

    for n, cruda in enumerate(crudas, start=2):
        def celda(clave):
            i = columnas.get(clave)
            return _texto(cruda[i]) if i is not None and i < len(cruda) else ""

        def descripcion_cruda():
            i = columnas.get("descripcion")
            return _texto_con_renglones(cruda[i]) if i is not None and i < len(cruda) else ""

        if not any(_texto(v) for v in cruda):
            continue
        if len(rev.filas) >= MAX_FILAS:
            rev.errores_generales.append(f"El archivo pasa de {MAX_FILAS} filas: divídalo en dos.")
            break
        f = Fila(numero=n, codigo_barras=celda("codigo_barras"), nombre=celda("nombre")[:200],
                 marca=celda("marca")[:80], presentacion=celda("presentacion")[:120],
                 sku=celda("sku")[:30], descripcion=descripcion_cruda(),
                 cabys="".join(celda("cabys").split()))
        if f.nombre.lower().startswith("ejemplo:"):
            continue   # la fila de ejemplo de la plantilla
        if f.cabys and not (f.cabys.isdigit() and len(f.cabys) == 13):
            f.errores.append(f"El CABYS «{f.cabys}» no es válido: son 13 dígitos.")

        for clave in ("cantidad", "costo_unitario", "precio_venta", "bonificadas"):
            valor = _numero(celda(clave))
            if valor is None:
                f.errores.append(f"«{celda(clave)}» no es un número ({clave.replace('_', ' ')}).")
                valor = Decimal("0")
            setattr(f, clave, str(valor) if (clave != "precio_venta" or valor > 0) else "")
        # 0 facturadas vale si la fila trae bonificadas: regalo del proveedor.
        if Decimal(f.cantidad) < 0 or (Decimal(f.cantidad) == 0 and Decimal(f.bonificadas) <= 0):
            f.errores.append("La cantidad tiene que ser mayor que cero (o la fila tiene que traer bonificadas).")
        if Decimal(f.costo_unitario) < 0:
            f.errores.append("El costo no puede ser negativo.")
        if Decimal(f.bonificadas) < 0:
            f.errores.append("Las bonificadas no pueden ser negativas.")

        # ¿Qué producto es? Primero por SKU, si la fila lo trae: es el código
        # que menos se repite y el que identifica al producto en todo el ERP.
        producto = None
        if f.sku:
            if f.sku in skus_vistos:
                f.errores.append(f"El código {f.sku} ya viene en la fila {skus_vistos[f.sku]}: "
                                 "sume las cantidades en una sola fila.")
            skus_vistos[f.sku] = n
            producto = por_sku.get(f.sku)
            if producto is None:
                if f.sku in skus_ocupados:
                    f.errores.append(f"El código {f.sku} ya lo usa un producto inactivo o de otra empresa. "
                                     "Reactívelo en el panel de administración o use otro código.")
                otro = por_codigo.get(f.codigo_barras) if f.codigo_barras else None
                if otro is not None:
                    # Crearlo duplicaría el producto con el mismo código de
                    # barras: la pistola ya no sabría cuál de los dos es.
                    f.errores.append(f"El código de barras {f.codigo_barras} ya es del producto "
                                     f"«{otro.nombre}» (código {otro.sku}). Revise cuál de los dos es el correcto.")
                elif f.codigo_barras and f.codigo_barras in codigos_ocupados:
                    f.errores.append(f"El código de barras {f.codigo_barras} ya lo usa un producto inactivo.")
                if f.codigo_barras and f.codigo_barras in codigos_vistos:
                    f.errores.append(f"El código de barras {f.codigo_barras} ya viene en la fila "
                                     f"{codigos_vistos[f.codigo_barras]}.")
                if f.codigo_barras:
                    codigos_vistos[f.codigo_barras] = n
        elif f.codigo_barras:
            producto = por_codigo.get(f.codigo_barras) or por_sku.get(f.codigo_barras)
            if f.codigo_barras in codigos_vistos:
                f.errores.append(f"El código {f.codigo_barras} ya viene en la fila {codigos_vistos[f.codigo_barras]}: "
                                 "sume las cantidades en una sola fila.")
            codigos_vistos[f.codigo_barras] = n
        # Con SKU no se busca por nombre: un código que no existe es la orden
        # de crear ESE producto, aunque ya haya otro que se llame igual (las
        # variantes de Buen Amigo comparten nombre y se distinguen por código).
        if producto is None and f.nombre and not f.codigo_barras and not f.sku:
            candidatos = por_nombre.get(_plano(f.nombre), [])
            if len(candidatos) > 1:
                f.errores.append(f"Hay {len(candidatos)} productos que se llaman «{f.nombre}»: "
                                 "ponga el código de barras para saber cuál es.")
            elif candidatos:
                producto = candidatos[0]
        if producto is not None:
            f.producto_id = producto.id
            f.nombre = producto.nombre
            f.precio_actual = str(producto.precio_venta)
        else:
            if not f.nombre:
                f.errores.append("No existe un producto con ese código: para crearlo escriba el nombre.")
            if not f.precio_venta:
                f.errores.append("Producto nuevo: falta el precio de venta.")

        categoria = producto.categoria if producto is not None else None
        if celda("categoria"):
            cat = buscar_categoria(celda("categoria"))
            if cat is None:
                f.errores.append(f"La categoría «{celda('categoria')}» no existe en el ERP.")
            else:
                f.categoria_id, f.categoria_txt = cat.id, str(cat)
                categoria = categoria or cat   # en un existente no se cambia (ver `aplicar`)

        # La foto del .zip: por el SKU o el código de barras de la fila, o
        # por el SKU del producto que se encontró.
        for clave in (f.sku, f.codigo_barras, producto.sku if producto is not None else ""):
            if clave and clave.lower() in fotos:
                f.foto_clave = clave.lower()
                fotos_usadas.add(f.foto_clave)
                if producto is not None and producto.imagen:
                    f.avisos.append("Ya tiene foto: la del .zip no se usa. Para cambiarla, Ajustes de inventario.")
                else:
                    f.foto_se_pone = True
                break
        raiz = (categoria.padre or categoria) if categoria is not None else None
        calculo.append((f, producto.tarifa_iva if producto is not None else TARIFA_GENERAL_IVA,
                        raiz is not None and raiz.nombre in RAICES_ALIMENTO))
        if celda("mascota"):
            f.mascota = normalizar_mascota(celda("mascota")) or ""
            if not f.mascota:
                f.errores.append(f"Mascota «{celda('mascota')}» no válida. Opciones: {', '.join(MASCOTAS)}.")
        rev.filas.append(f)

    if not rev.filas and not rev.errores_generales:
        rev.errores_generales.append("El archivo no trae ninguna fila con mercadería.")
    sobran = sorted(fotos[c][0] for c in set(fotos) - fotos_usadas)
    if sobran:
        rev.errores_generales.append(
            f"Estas fotos no corresponden a ninguna fila (el nombre tiene que ser el código de la fila): "
            f"{', '.join(sobran[:15])}{'…' if len(sobran) > 15 else ''}.")
    if rev.descuento < 0:
        rev.errores_generales.append("El descuento de la factura no puede ser negativo.")
    elif rev.descuento and rev.descuento >= rev.total:
        rev.errores_generales.append(f"El descuento (₡{rev.descuento}) no puede ser igual o mayor que el total "
                                     f"de la factura (₡{rev.total}).")
    else:
        tradicional = empresa.regimen == Empresa.Regimen.TRADICIONAL
        for f, tarifa, alimento in calculo:
            _calcular_margen(f, rev, tarifa if tradicional else Decimal("0"), alimento)
    return rev


def costo_real(f, rev) -> Decimal:
    """Lo que de verdad cuesta cada unidad que entra a bodega: el total de la
    línea menos su parte del descuento de la factura, repartido entre las
    facturadas y las de regalo. Mismo criterio que `LineaCompra.costo_real_unitario`."""
    unidades = Decimal(f.cantidad) + Decimal(f.bonificadas)
    if unidades <= 0:
        return Decimal("0")
    total = f.total
    if rev.descuento and rev.total:
        total -= rev.descuento * total / rev.total
    return total / unidades


def _calcular_margen(f, rev, tarifa, alimento):
    precio = Decimal(f.precio_venta or f.precio_actual or "0")
    if precio <= 0 or f.errores:
        return
    sin_iva = precio / (1 + tarifa / 100)
    costo = costo_real(f, rev)
    margen = (sin_iva - costo) / sin_iva * 100
    minimo = MARGEN_MINIMO_ALIMENTO if alimento else MARGEN_MINIMO_OTROS
    f.margen = str(margen.quantize(Decimal("0.1"), ROUND_HALF_UP))
    f.margen_minimo = str(minimo)
    if margen < minimo:
        # El precio que deja justo el mínimo, con IVA y en colones enteros hacia arriba.
        precio_min = (costo / (1 - minimo / 100) * (1 + tarifa / 100)).quantize(Decimal("1"), "ROUND_CEILING")
        precio_txt = f"{precio_min:,.0f}".replace(",", ".")
        f.avisos.append(f"Margen {f.margen.replace('.', ',')} %, debajo del mínimo de {minimo:g} %. "
                        f"Con el mínimo el precio sería ₡{precio_txt}.")


def a_datos(rev: Revision) -> list[dict]:
    return [asdict(f) for f in rev.filas]


@transaction.atomic
def aplicar(*, empresa, filas: list[dict], proveedor, sucursal, forma_pago, factura_proveedor, iva, usuario,
            descuento=0):
    """Crea los productos nuevos, registra y recibe la compra, cambia los
    precios y pone las fotos. Todo o nada en la base. Devuelve la compra.

    Las fotos se copian al bucket dentro de la transacción: si algo falla
    después, el producto no queda apuntando a nada, pero sí puede quedar un
    archivo suelto en el bucket — no estorba y es preferible a un producto
    sin foto."""
    lineas, precios, con_foto = [], [], []
    for datos in filas:
        f = Fila(**{k: v for k, v in datos.items() if k in Fila.__dataclass_fields__})
        if f.producto_id:
            producto = Producto.objects.get(pk=f.producto_id, empresa=empresa)
            if f.cambia_precio:
                precios.append((producto, Decimal(f.precio_venta)))
            campos = []
            for campo in ("marca", "presentacion", "mascota", "descripcion", "cabys"):
                if getattr(f, campo) and not getattr(producto, campo):
                    setattr(producto, campo, getattr(f, campo))
                    campos.append(campo)
            if f.categoria_id and not producto.categoria_id:
                producto.categoria_id = f.categoria_id
                campos.append("categoria")
            # Encontrado por SKU y sin código de barras propio: se le pone el
            # de la fila (el de fábrica), salvo que ya sea de otro producto.
            if (f.sku and f.codigo_barras and not producto.codigo_barras
                    and not Producto.objects.filter(codigo_barras=f.codigo_barras).exists()):
                producto.codigo_barras = f.codigo_barras
                campos.append("codigo_barras")
            if campos:   # solo completa lo vacío: nunca pisa datos puestos a mano
                producto.save(update_fields=campos + ["actualizado_en"])
        else:
            if f.codigo_barras and Producto.objects.filter(codigo_barras=f.codigo_barras).exists():
                raise ValidationError(f"Fila {f.numero}: el código {f.codigo_barras} ya existe.")
            if f.sku and Producto.objects.filter(sku=f.sku).exists():
                raise ValidationError(f"Fila {f.numero}: el código {f.sku} ya existe.")
            from .views import _generar_sku
            producto = Producto.objects.create(
                empresa=empresa, sku=f.sku or _generar_sku(), nombre=f.nombre, codigo_barras=f.codigo_barras,
                precio_venta=Decimal(f.precio_venta), marca=f.marca, presentacion=f.presentacion,
                mascota=f.mascota, categoria_id=f.categoria_id, descripcion=f.descripcion, cabys=f.cabys,
            )
        if f.foto_se_pone and f.foto and not producto.imagen:
            con_foto.append((producto, f.foto))
        lineas.append({"producto": producto, "cantidad": Decimal(f.cantidad),
                       "cantidad_bonificada": Decimal(f.bonificadas), "costo_unitario": Decimal(f.costo_unitario)})
    compra = services.crear_y_recibir_compra(
        proveedor=proveedor, sucursal=sucursal, lineas=lineas, forma_pago=forma_pago,
        factura_proveedor=factura_proveedor, usuario=usuario, iva=iva, descuento=descuento,
    )
    sello = f"{timezone.now():%Y%m%d%H%M%S}"
    for producto, ruta in con_foto:
        with default_storage.open(ruta, "rb") as archivo:
            contenido = archivo.read()
        # Con la fecha en el nombre: el sitio y Cloudflare guardan cada
        # dirección de foto hasta 30 días (mismo criterio que Ajustes).
        producto.imagen = guardar_imagen_producto(f"{producto.sku}-{sello}{posixpath.splitext(ruta)[1]}", contenido)
        producto.save(update_fields=["imagen", "actualizado_en"])
    for producto, precio in precios:
        cambiar_precio(producto=producto, nuevo_precio=precio, usuario=usuario,
                       motivo=f"Carga masiva {compra.numero}")
    return compra
